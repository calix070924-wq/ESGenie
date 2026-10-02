"""R8p~r·R8o 웹 변형 출력 대조(review-r5/web_variant_checks.py와 같은 도구, R8r의 377,000 MJ 계획 행 숫자와 행 이름·근거 성격 기록 추가) — 화면 상세·Excel·번들 PDF의 수치 등장 횟수, 최종 답변 근거 bbox가 PDF의 어느 칸 글자인지.
칸 위치는 변형 PDF 자체에서 fitz search_for로 구한다(제품 코드와 독립). 사용: web_variant_checks.py <review-r6> <sha7> <out.json> [변형 목록 R8p,R8q] [서버·화면 폴더 접미사 예: _r6before]
R8o 입력 PDF는 review-r5/variant_inputs/에 있다(다시 만들지 않음)."""
import json
import re
import sys
import zipfile
from pathlib import Path

import fitz
import openpyxl

r4, sha, out = Path(sys.argv[1]), sys.argv[2], Path(sys.argv[3])
names = sys.argv[4].split(",") if len(sys.argv) > 4 else ["R8p", "R8q", "R8r", "R8o"]
tag = sys.argv[5] if len(sys.argv) > 5 else ""
NEEDLES = ["247,500", "247500", "0.2475", "13.885", "360,772", "360772", "0.360772", "20.239", "8,420", "8420",
           "377,000", "377000", "0.377", "21.15", "8,800", "8800"]
runs = {json.loads((d / "input_project.json").read_text(encoding="utf-8"))["id"]: d
        for d in (r4 / "live" / f"web_server_{sha}{tag}" / "runs").glob("*_rev*") if (d / "input_project.json").exists()}


def count(text):
    t = re.sub(r"\s+", "", text)
    return {n: t.count(n) for n in NEEDLES}


def cell_boxes(pdf):
    with fitz.open(pdf) as doc:
        page = doc[0]
        w, h = page.rect.width, page.rect.height
        return {s: [[round(r.x0 / w, 4), round(r.y0 / h, 4), round(r.x1 / w, 4), round(r.y1 / h, 4)] for r in page.search_for(s)]
                for s in ("360,772", "247,500", "8,420", "377,000", "8,800")}


def which_cell(bbox, boxes):
    cx, cy = (bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2
    hits = [s for s, rs in boxes.items() for b in rs if b[0] - 0.01 <= cx <= b[2] + 0.01 and b[1] - 0.01 <= cy <= b[3] + 0.01]
    return hits or ["(표 칸 글자 밖)"]


res = {}
for name in names:
    vd = r4 / "live" / f"web_variant_{name}_{sha}"
    pid = json.loads((vd / "session.json").read_text(encoding="utf-8"))["project_id"]
    run = runs[pid]
    pdf = next(p for d in (r4 / "variant_inputs", r4.parent / "review-r5" / "variant_inputs") for p in d.glob(f"{name}_*.pdf"))
    boxes = cell_boxes(pdf)
    pipe = json.loads((run / "pipeline.json").read_text(encoding="utf-8"))
    result = json.loads((run / "result.json").read_text(encoding="utf-8"))
    ext = pipe["ocr_extractions"][0]
    answers = {}
    for a in result["sheet"]["answers"]:
        code = a["qid"][-5:]
        if code in ("E-4-1", "E-3-1"):
            answers[code] = {"value": a.get("value"), "unit": a.get("unit"), "status": a.get("status"),
                             "evidence": [{"quote": l["quote"], "bbox": l.get("bbox"), "cell": which_cell(l["bbox"], boxes) if l.get("bbox") else None}
                                          for l in a.get("evidence_links", [])],
                             "reference": [{"quote": l["quote"], "bbox": l.get("bbox")} for l in a.get("reference_links", [])]}
    ui = json.loads((r4 / "live" / f"ui_screens_{sha}{tag}" / name / "detail_text.json").read_text(encoding="utf-8"))
    ui_text = "".join(v["row"] + v["detail"] for v in ui.values())
    wb = openpyxl.load_workbook(vd / "검토용초안.xlsx")
    xlsx_text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)
    xlsx_rows = {str(row[0].value): [str(c.value) for c in row if c.value is not None]
                 for ws in wb.worksheets for row in ws.iter_rows() if row and str(row[0].value or "").endswith(("E-4-1", "E-3-1"))}
    res[name] = {
        "pdf": pdf.name, "pdf_sha256": ext["router_meta"].get("source_sha256"), "project_id": pid, "run_dir": run.name,
        "engine": ext["router_meta"].get("engine"), "upstage_model": ext["router_meta"].get("upstage_model"),
        "metrics": [(m["value"], m["unit"], m.get("kesg_code_guess")) for m in ext["metrics"]],
        "metric_rows": [{"value": m["value"], "unit": m["unit"], "hint": m.get("metric_hint"),
                         "row_label": (m.get("source_detail") or {}).get("row_label"),
                         "row_item": (m.get("source_detail") or {}).get("row_item")} for m in ext["metrics"]],
        "graph_basis": [(n["metric"], n["value"], n["unit"], (n.get("boundary") or {}).get("basis")) for n in pipe["evidence_graph"]["nodes"]],
        "table_review": [(r.get("reason"), r.get("row_label")) for r in (ext["router_meta"].get("table_metrics") or {}).get("review", [])],
        "cell_boxes_from_pdf": boxes, "answers": answers, "ui_rows": {k: v["row"] for k, v in ui.items()},
        "ui": count(ui_text), "pdf_bundle": count((vd / "report_text.txt").read_text(encoding="utf-8")), "xlsx": count(xlsx_text),
        "xlsx_rows": xlsx_rows,
    }
out.write_text(json.dumps({"code_sha": sha, "note": "공백 제거 후 부분 문자열 개수. ui=화면 목록·상세(E-4-1·E-3-1, details 펼침), "
                           "pdf_bundle=번들 실사응답서 본문, xlsx=전체 시트 셀. cell=근거 bbox 중심이 변형 PDF의 어느 숫자 글자 상자 안인지",
                           "results": res}, ensure_ascii=False, indent=1), encoding="utf-8")
for n, r in res.items():
    print(n, r["engine"], r["metrics"], r["table_review"], [(m["row_label"], m["row_item"]) for m in r["metric_rows"]], r["graph_basis"])
    for c, a in r["answers"].items():
        print("  ", c, a["value"], a["unit"], a["status"], [e["cell"] for e in a["evidence"]], len(a["reference"]))
    print("   fee counts ui/pdf/xlsx", [r[k]["247,500"] + r[k]["247500"] + r[k]["0.2475"] + r[k]["13.885"] for k in ("ui", "pdf_bundle", "xlsx")],
          "heat", [r[k]["360772"] + r[k]["360,772"] for k in ("ui", "pdf_bundle", "xlsx")], "plan377",
          [r[k]["377,000"] + r[k]["377000"] + r[k]["0.377"] + r[k]["21.15"] for k in ("ui", "pdf_bundle", "xlsx")], "TJ/t", [(r[k]["0.360772"], r[k]["20.239"]) for k in ("ui", "pdf_bundle", "xlsx")])
