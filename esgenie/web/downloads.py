"""Exports include both engine provenance and separately labeled human notes."""
from __future__ import annotations

from copy import deepcopy
from dataclasses import fields
from io import BytesIO
import json
from pathlib import Path
import re
from tempfile import TemporaryDirectory
from zipfile import ZipFile, ZIP_DEFLATED

from .store import WorkspaceError


def build_download(project: dict, directory: Path, kind: str) -> tuple[bytes, str, str]:
    from esgenie.supplychain.schema import Answer, ResponseSheet
    from esgenie.ssot.audit_trace import EvidenceLink
    from esgenie.supplychain import copy_evidence_pack, export_response_sheet, export_response_sheet_pdf
    from openpyxl import load_workbook

    result = project.get("result")
    if not result:
        raise WorkspaceError("서류를 읽은 뒤 응답서를 받을 수 있어요.", 409)
    if project["input_revision"] != project["result_revision"]:
        raise WorkspaceError("자료가 바뀌었어요. 다시 분석한 뒤 최신 응답서를 받아 주세요.", 409)
    if project["job"]["status"] in {"queued", "running"}:
        raise WorkspaceError("자료 읽기가 끝난 뒤 응답서를 받아 주세요.", 409)
    raw = deepcopy(result["sheet"])
    answer_fields = {f.name for f in fields(Answer)}
    link_fields = {f.name for f in fields(EvidenceLink)}
    answers = []
    presented = {a["id"]: a for a in result["answers"]}
    for item in raw["answers"]:
        item = {k: v for k, v in item.items() if k in answer_fields}
        item["evidence_links"] = [EvidenceLink(**{k: v for k, v in link.items() if k in link_fields}) for link in item.get("evidence_links", [])]
        item["flags"] = list(item.get("flags", [])) + presented[item["qid"]]["notices"]
        answers.append(Answer(**item))
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", project["company_name"])[:80]
    prefix = "사용법 예시" if project["mode"] == "example" else "검토용 초안"
    sheet = ResponseSheet(raw["framework_key"], f"[{prefix}] {raw['framework_label']}", name, answers, raw.get("gaps", []))
    notes = [f"# {prefix} · {project['company_name']}", "", "직접 작성한 답변과 메모는 검증된 사실로 자동 변경되지 않습니다.", ""]
    for answer in answers:
        note = project["notes"].get(answer.qid)
        if note:
            old = note["revision"] != project["result_revision"]
            notes += [f"## {answer.question_text}", "", "이전 분석에서 작성한 기록입니다. 다시 확인해 주세요." if old else "담당자가 직접 작성한 내용 · 확인 전",
                      "", "직접 작성한 답변: " + note.get("answer", ""), "", "검토 메모: " + note.get("text", ""), ""]
    notes += ["## 추가 확인 사항", ""] + [f"- {message}" for message in result.get("limitations", [])]
    with TemporaryDirectory(prefix="download-", dir=directory) as temporary:
        out = Path(temporary)
        excel_path = Path(export_response_sheet(sheet, out))
        workbook = load_workbook(excel_path)
        review = workbook.create_sheet("담당자 작성")
        review.append(["질문", "직접 작성한 답변 (확인 전)", "검토 메모", "기록 기준"])
        for answer in answers:
            note = project["notes"].get(answer.qid)
            if note:
                review.append([answer.question_text, note.get("answer", ""), note.get("text", ""),
                               "이전 분석 기록 · 재확인 필요" if note["revision"] != project["result_revision"] else "현재 분석"])
        review.column_dimensions["A"].width = 45
        review.column_dimensions["B"].width = 55
        review.column_dimensions["C"].width = 55
        # Uploaded text and human answers are text, never spreadsheet formulas.
        for worksheet in workbook:
            for row in worksheet:
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                        cell.value = "'" + cell.value
        workbook.save(excel_path)
        if kind == "xlsx":
            return excel_path.read_bytes(), f"{prefix}_{name}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        paths = {doc["name"]: str(directory / "uploads" / doc["id"] / doc["name"])
                 for doc in project["documents"] if not doc.get("example")}
        copy_evidence_pack(sheet, out, paths)
        export_response_sheet_pdf(sheet, out, evidence_base_dir=out, embed_evidence=project["mode"] != "example")
        (out / "담당자_작성_및_확인사항.md").write_text("\n".join(notes), encoding="utf-8")
        (out / "검증_근거.json").write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
        archive = BytesIO()
        with ZipFile(archive, "w", ZIP_DEFLATED) as bundle:
            for path in sorted(out.rglob("*")):
                if path.is_file():
                    bundle.write(path, path.relative_to(out).as_posix())
        return archive.getvalue(), f"{prefix}_{name}_제출준비.zip", "application/zip"
