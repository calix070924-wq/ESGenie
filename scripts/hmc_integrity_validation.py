"""실제 시연 PDF/저장 OCR → 원장 → D1 → 답변 → 3종 출력의 필수 검증.

네트워크/라이브 LLM 없음. --evidence-dir, --cache-dir는 원본을 읽기만 한다.
수치/상태를 Answer나 DataPoint에 주입하지 않는다. 동일 범위 폐기물 92% 주장은
명시적으로 분리한 통제 실험이다. 생성 LLM만 결정적 원문 발췌로 대체하며 실제
적합성·grounding 게이트는 실행한다. 필수 입력 또는 검사가 없으면 실패한다.
"""
from __future__ import annotations
import argparse
import copy
import hashlib
import json
import os
import re
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

os.environ.setdefault("ESGENIE_FORCE_MOCK", "1")
REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
OUT_DIR = REPO / "outputs/validation/hmc-response-integrity-followup-20260921"
LOG_DIR = REPO / "docs/validation/hmc-response-integrity-followup-20260921"
CACHE_NAME = "0ad49a4822ee713d7a36ba8cce53c0e31e9c47cf1eaf0f1cdb80cc137f318ce9.json"
CLAUSE_DOCS = (
    "의사소통절차서_2025.pdf", "15_노사협의회규정_2026-06.pdf",
    "17_고충처리및신원보호규정_2026-06.pdf", "18_정보보호정책_2026-06.pdf",
    "16_부패방지윤리경영규정_2026-06.pdf", "14_임금및복리후생규정_2026-06.pdf",
    "근로시간관리규정_2025.pdf", "지식재산보호규정_2025.pdf",
    "위생_기숙사관리규정_2025.pdf", "책임광물실사정책_2025.pdf",
)
# 이 검증기의 명시적 인수 조건이다. 제품의 계산/판정에는 전달하지 않는다.
REQUIRED = {
    "energy": "HMC-C-8-E-4-1", "renewable": "HMC-C-8-E-4-2",
    "emissions": "HMC-C-8-E-3-1", "waste": "HMC-C-4-E-6-2", "communication": "HMC-E-7",
}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def dump(path, value):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def pdf_pages(path):
    import fitz
    with fitz.open(path) as doc:
        return [p.get_text() for p in doc]


def pdf_text(path):
    return "\n".join(pdf_pages(path))


def hmc_fitness():
    from esgenie.knowledge.kesg_evidence_requirements import requirement_for_question
    from esgenie.supplychain import get_framework
    from esgenie.supplychain.question_fitness import build_fitness_map
    fw = get_framework("hmc")
    return fw, build_fitness_map(fw.questions, lambda q: requirement_for_question(q.kesg_codes, quantitative=q.qtype == "numeric"))


def load_actual_inputs(evidence_dir, cache_dir):
    from esgenie.ssot.ocr_router import (
        _extract_structured_no_llm, _map_vlm_json, _backfill_kesg_codes,
        OcrExtraction, DocChannel, ExtractedClause,
    )
    inputs, provenance = [], []
    for name, kind in (
        ("01_전기요금청구서_2026-05.pdf", "kepco_bill"),
        ("02_도시가스요금고지서_2026-05.pdf", "gas_bill"),
        ("03_사업장폐기물_위탁처리명세_2026-04.pdf", "waste_ledger"),
    ):
        path = evidence_dir / name
        assert path.is_file(), f"필수 원문 없음: {path}"
        ext = _extract_structured_no_llm(str(path), doc_type=kind)
        assert ext.metrics and ext.raw_text, f"실제 로컬 추출 실패: {name}"
        ext.router_meta.update(source_sha256=sha256(path), validation_input="local_pdf_parser")
        inputs.append(ext)
        provenance.append({"path": str(path), "sha256": sha256(path), "pages_zero_based": list(range(len(pdf_pages(path)))), "input_class": "local PDF text + production structured parser", "metrics": [m.__dict__ for m in ext.metrics]})
    name = "12_재생에너지사용현황_2026-06.pdf"
    path, cache_path = evidence_dir / name, cache_dir / CACHE_NAME
    assert path.is_file() and cache_path.is_file(), "필수 12번 원문/저장 OCR 캐시 없음"
    saved = json.loads(cache_path.read_text())
    assert saved["meta"]["source_file"] == name, "다른 문서의 캐시"
    raw = pdf_text(path)
    metrics, clauses = _map_vlm_json(saved["response"], page_no=1)
    # 캐시에는 입력 원문 해시가 없다. 파일명 메타데이터 및 실제 값·어휘를 확인하고
    # 이 연결의 한계를 보고한다. 원문 텍스트는 문맥/모순 검출에 별도로 공급한다.
    assert len(pdf_pages(path)) == 1 and "10.6" in raw and "태양광" in raw
    assert any(m.value == 10.6 and m.unit == "%" for m in metrics)
    ext = OcrExtraction(name, DocChannel.UNSTRUCTURED, "saved_vlm_replay", metrics=metrics, clauses=clauses, raw_text=raw,
        router_meta={"source_sha256": sha256(path), "cache_sha256": sha256(cache_path), "validation_input": "saved OCR + original text context"})
    _backfill_kesg_codes(ext)
    inputs.append(ext)
    provenance.append({"path": str(path), "sha256": sha256(path), "pages_zero_based": [0], "cache_path": str(cache_path), "cache_sha256": sha256(cache_path), "cache_meta": saved["meta"], "input_class": "saved OCR response replay; context augmented from original PDF", "link_limit": "cache metadata lacks original input SHA256; filename and source content checked"})
    fw, _ = hmc_fitness()
    code = next(q.primary_code for q in fw.questions if q.qid == REQUIRED["communication"])
    path = evidence_dir / CLAUSE_DOCS[0]
    # 페이지 위치는 원문 PDF의 enumerate에서만 얻는다.
    pages = pdf_pages(path)
    ext = OcrExtraction(path.name, DocChannel.UNSTRUCTURED, "local_clause_text",
        clauses=[ExtractedClause(section="의사소통", text=t, kesg_code_guess=code, page=i) for i, t in enumerate(pages)],
        raw_text="\n".join(pages), router_meta={"source_sha256": sha256(path), "validation_input": "local_pdf_clause"})
    inputs.append(ext)
    provenance.append({"path": str(path), "sha256": sha256(path), "pages_zero_based": list(range(len(pages))), "input_class": "local PDF text; question code manually routed, fitness and grounding executed"})
    return inputs, provenance


def pipeline(inputs, claims=None, *, label="실제자료"):
    from esgenie.ssot.evidence_graph import build_unified_graph
    from esgenie.ssot.ssot_pipeline import extract_with_ssot
    from esgenie.ssot.detector_5axis import detect_d1_numeric
    from esgenie.ssot.audit_trace import build_data_points
    from esgenie.supplychain.responder import build_response_sheet
    graph = build_unified_graph(None, inputs, corp_code="HMC_REPLAY", corp_name="한울정밀공업", report_year=2026)
    report = SimpleNamespace(source="ssot_local", corp_code="HMC_REPLAY", corp_name="한울정밀공업", report_year=2026, fiscal_year=2026, kesg_data={}, sections={}, raw_text="")
    extraction = extract_with_ssot(report, graph, profile="sme")
    codes = [c for c, f in graph.resolved_facts.items() if f]
    scores, evaluations = {}, {}
    for code in codes:
        fact = graph.resolved_facts[code]
        axis = detect_d1_numeric(f"{extraction.mapped[code]['name']} {fact.value}{fact.unit}", code, graph)
        scores[code], evaluations[code] = axis.score, axis.evaluation
    points = build_data_points(graph, scores, target_codes=codes, d1_evaluations=evaluations)
    sheet = build_response_sheet("hmc", corp_name=f"한울정밀공업_{label}", extraction=extraction, data_points=points, supplier_claims=claims)
    return graph, extraction, points, sheet, evaluations


def draft_actual_communication(sheet, graph):
    """생성 LLM만 결정적 발췌 대체. 실제 grounding/적합성/승인대기 경로 실행."""
    from esgenie.supplychain.drafter import generate_drafts
    node = next(n for n in graph.text_nodes.values() if n.source_file == CLAUSE_DOCS[0])
    sentence = "ESG 방침·관행·기대·성과를 근로자·공급사·고객에 명확하고 정확하게 전달하는 프로세스를 규정한다."
    assert re.sub(r"\s+", "", sentence) in re.sub(r"\s+", "", node.text), "원문에 없는 결정적 초안"
    llm = MagicMock()
    llm.complete.side_effect = lambda **kw: SimpleNamespace(content=(f"{sentence} [{node.id}]" if kw.get("mock_hint") == "generate" and "[E-7]" in kw.get("user", "") else "INSUFFICIENT_EVIDENCE"))
    with patch("esgenie.supplychain.drafter.LLMClient", return_value=llm):
        sheet = generate_drafts(sheet, graph)
    return sheet


def sheet_payload(sheet):
    from esgenie.supplychain.checklist import build_checklist
    return {"framework_key": sheet.framework_key, "framework_label": sheet.framework_label, "corp_name": sheet.corp_name,
        "summary": {k: getattr(sheet, k) for k in ("auto_pct", "draft_pct", "hitl_pct", "pending_pct", "flagged_count")},
        "answers": [a.to_dict() for a in sheet.answers], "checklist": [x.to_dict() for x in build_checklist(sheet)]}


def validate_payload(payload, *, controlled):
    """저장된 JSON 자체를 검사. 필수 qid별로 비어 있지 않은 조건을 요구한다."""
    failures, checks = [], []
    def check(name, ok):
        checks.append({"check": name, "pass": bool(ok)})
        if not ok:
            failures.append(name)
    answers = {a["qid"]: a for a in payload.get("answers", [])}
    checklist = {a["qid"]: a for a in payload.get("checklist", [])}
    from esgenie.supplychain import get_framework
    expected_ids = {q.qid for q in get_framework("hmc").questions}
    check("all_question_ids", set(answers) == expected_ids and len(payload.get("answers", [])) == len(expected_ids))
    for name, qid in REQUIRED.items():
        check(f"{name}:required", qid in answers)
        if qid not in answers:
            continue
        a = answers[qid]
        if name == "communication":
            check("communication:approval_pending", a["status"] == "draft_ready")
            cites = a.get("draft_citations", [])
            check("communication:actual_page", bool(cites) and all(c["source_file"] == CLAUSE_DOCS[0] and c.get("page") == 0 for c in cites))
            check("communication:grounding", bool(a.get("draft_grounding")) and a["draft_grounding"].get("decision") == "ACCEPT" and not a["draft_grounding"].get("soft_flags"))
            continue
        expected, unit = {"energy": (.873988, "TJ"), "renewable": (10.6, "%"), "emissions": (88.397, "tCO2eq"), "waste": (29.3, "%")}[name]
        check(f"{name}:value_unit", a.get("value") == expected and a.get("unit") == unit)
        check(f"{name}:boundary", bool(a.get("boundary_label")) and bool(a.get("boundary")))
        check(f"{name}:source", bool(a.get("evidence_links")))
        expected_sources = ({"01_전기요금청구서_2026-05.pdf", "02_도시가스요금고지서_2026-05.pdf"}
            if name in ("energy", "emissions") else
            {"12_재생에너지사용현황_2026-06.pdf"} if name == "renewable" else
            {"03_사업장폐기물_위탁처리명세_2026-04.pdf"})
        check(f"{name}:source_identity_page", {e.get("file_name") for e in a.get("evidence_links", [])} == expected_sources
              and all(e.get("page") == 0 for e in a.get("evidence_links", [])))
        b = a.get("boundary", {})
        if name in ("energy", "emissions"):
            check(f"{name}:period_site", (b.get("period_start"), b.get("period_end")) == ("2026-04-25", "2026-05-24") and b.get("site_scope") == "site")
            check(f"{name}:two_sources_partial", len(a.get("evidence_links", [])) == 2 and a.get("completeness") == "partial" and a.get("status") != "verified")
        if name in ("energy", "renewable"):
            check(f"{name}:scope_action", qid in checklist and checklist[qid].get("action") == "범위 확인·보완" and bool(checklist[qid].get("request")))
        if name == "renewable":
            notes = " ".join(a.get("scope_notes", []))
            check("renewable:conflict_and_denominator", "설비 동일성" in notes and "분모" in notes and "조달수단" in notes and a["status"] != "verified")
        if name == "waste":
            check("waste:period", b.get("period_start") == "2026-04-01" and b.get("period_end") == "2026-04-30")
            check("waste:comparison", a.get("comparison") == ("mismatch" if controlled else "scope_unconfirmed"))
            if controlled:
                check("waste:mismatch_action", qid in checklist and "92" in checklist[qid].get("request", "") and "29.3" in checklist[qid].get("request", ""))
    check("summary:four_parts", abs(sum(payload.get("summary", {}).get(k, 0) for k in ("auto_pct", "draft_pct", "hitl_pct", "pending_pct"))-100) < .2)
    return {"checks": checks, "failed": failures, "executed": len(checks)}


def negative_controls(payload):
    variants = {}
    def mutate(label, fn):
        p = copy.deepcopy(payload); fn(p)
        result = validate_payload(p, controlled=True)
        assert result["failed"], f"음성 대조군이 통과함: {label}"
        variants[label] = result["failed"]
    qid = REQUIRED["energy"]
    get = lambda p, q=qid: next(a for a in p["answers"] if a["qid"] == q)
    mutate("missing_question", lambda p: p["answers"].remove(get(p, REQUIRED["waste"])))
    mutate("empty_boundary", lambda p: get(p).update(boundary={}, boundary_label=""))
    mutate("changed_value", lambda p: get(p).update(value=17.6))
    mutate("wrong_page", lambda p: get(p, REQUIRED["communication"])["draft_citations"][0].update(page=1))
    mutate("missing_checklist", lambda p: p.update(checklist=[x for x in p["checklist"] if x["qid"] != qid]))
    return variants


def normalize(text):
    # PDF 줄바꿈과 폰트 기호 변환만 정규화한다. 수치·단위를 보존한다.
    from esgenie.supplychain.exporters._fonts import pdf_safe_text
    return re.sub(r"\s+", "", pdf_safe_text(str(text or "")))


def pdf_answer_regions(path):
    """PDF 표의 가로선에서 행 영역을 구해 qid와 같은 행만 대조한다."""
    import fitz
    regions = {}
    checklist_started = False
    with fitz.open(path) as doc:
        for page_index, page in enumerate(doc):
            if checklist_started:
                break
            checklist_headers = page.search_for("제출 전 증빙 체크리스트")
            cutoff = checklist_headers[0].y0 if checklist_headers else page.rect.height
            lines = sorted({round(it[1].y, 2) for d in page.get_drawings() for it in d["items"] if it[0] == "l" and abs(it[1].y-it[2].y) < .1 and abs(it[2].x-it[1].x) > page.rect.width*.7})
            for top, bottom in zip(lines, lines[1:]):
                if bottom > cutoff:
                    continue
                rect = fitz.Rect(30, top, page.rect.width-30, bottom)
                left = page.get_text(clip=fitz.Rect(30, top, 99, bottom))
                qid = re.sub(r"\s+", "", left)
                if qid in set(REQUIRED.values()):
                    regions.setdefault(qid, []).append({"page": page_index+1, "rect": list(rect), "text": page.get_text(clip=rect)})
            checklist_started = bool(checklist_headers)
    return regions


def export_and_reopen(sheet, out_dir, *, controlled):
    from openpyxl import load_workbook
    from esgenie.supplychain.exporters.excel import export_response_sheet
    from esgenie.supplychain.exporters.pdf import export_response_sheet_pdf
    from esgenie.supplychain.render import draft_lines, note_lines, scope_line
    out_dir.mkdir(parents=True, exist_ok=True)
    json_path = out_dir / "response_sheet.json"
    dump(json_path, sheet_payload(sheet))
    xlsx = Path(export_response_sheet(sheet, out_dir))
    pdf = Path(export_response_sheet_pdf(sheet, out_dir, embed_evidence=False))
    payload = json.loads(json_path.read_text())
    result = validate_payload(payload, controlled=controlled)
    wb = load_workbook(xlsx)
    ws, cw = wb["응답서"], wb["증빙 체크리스트"]
    rows = {r[0].value: r for r in ws.iter_rows() if r[0].value}
    checklist_rows = {r[0].value: r for r in cw.iter_rows() if r[0].value}
    regions = pdf_answer_regions(pdf)
    def check(name, ok):
        result["checks"].append({"check": name, "pass": bool(ok)})
        if not ok: result["failed"].append(name)
    for qid in REQUIRED.values():
        a = next(a for a in sheet.answers if a.qid == qid)
        row = rows.get(qid)
        check(f"xlsx:{qid}:row", row is not None)
        if row:
            expected = ["\n".join(draft_lines(a)) if a.status == "draft_ready" else a.display_value, scope_line(a) or "—", a.badge, "\n".join(note_lines(a)).strip()]
            check(f"xlsx:{qid}:all_fields", [row[i].value or "" for i in (3,4,5,6)] == expected)
        if qid in {x["qid"] for x in payload["checklist"]}:
            check(f"xlsx:{qid}:checklist", qid in checklist_rows)
        check(f"pdf:{qid}:region", bool(regions.get(qid)))
        text = normalize(" ".join(x["text"] for x in regions.get(qid, [])))
        display = "\n".join(draft_lines(a)) if a.status == "draft_ready" else a.display_value
        check(f"pdf:{qid}:value_unit_or_draft", normalize(display) in text)
        check(f"pdf:{qid}:badge", normalize(a.badge) in text)
        check(f"pdf:{qid}:scope", not a.boundary_label or normalize(a.boundary_label) in text)
        check(f"pdf:{qid}:notes", all(normalize(n) in text for n in note_lines(a)))
    result.update(executed=len(result["checks"]), files={k: {"path": str(p), "sha256": sha256(p)} for k,p in (("json",json_path),("xlsx",xlsx),("pdf",pdf))}, pdf_regions=regions,
        xlsx_rows={qid: rows[qid][0].row for qid in REQUIRED.values() if qid in rows})
    return result


def replay_real_documents(evidence_dir):
    from esgenie.supplychain.question_fitness import select_fit_chunks
    from esgenie.rag_gates.grounding_gate import evaluate_grounding
    _, fmap = hmc_fitness()
    # 검토자가 판정한 관련도 라벨. 판정기에 주입하지 않으며 원문 발췌를 함께 기록한다.
    labels = {
        "HMC-E-7": ["관련", "무관", "무관", "무관", "무관", "무관", "무관", "무관", "무관", "무관"],
        "HMC-E-8": ["부분 관련", "관련", "관련", "무관", "부분 관련", "무관", "무관", "무관", "무관", "무관"],
    }
    texts = {n: pdf_text(evidence_dir/n) for n in CLAUSE_DOCS}  # 필수 입력이 없으면 예외
    rows, mixed = [], {}
    for i, (name,text) in enumerate(texts.items()):
        row = {"file": str(evidence_dir/name), "sha256": sha256(evidence_dir/name), "page": 0, "source_excerpt": text[:600]}
        for qid in labels:
            selected, why = select_fit_chunks([{"id":name,"text":text}], fmap[qid])
            expected = labels[qid][i] != "무관"
            assert bool(selected) == expected, f"수동 라벨과 단독 판정 상이: {qid}/{name}"
            row[qid] = {"manual_label": labels[qid][i], "selected": bool(selected), "hold_reason": why, "kept_excerpt": selected[0]["text"] if selected else ""}
        rows.append(row)
    for qid in labels:
        chunks = [{"id":n,"text":t} for n,t in texts.items()]
        expected = sorted(n for i,n in enumerate(CLAUSE_DOCS) if labels[qid][i] != "무관")
        orders = [chunks, chunks[::-1], chunks[3:]+chunks[:3]]
        results = [select_fit_chunks(order, fmap[qid])[0] for order in orders]
        assert all([x["id"] for x in result] == expected for result in results), f"혼합/순서 판정 상이:{qid}"
        mixed[qid] = {"kept": expected, "permutations": 3, "stable": True}
    source = [{"id":"SEC", "text":texts[CLAUSE_DOCS[3]]}]
    wrong = evaluate_grounding("ISMS 인증을 취득했다. [SEC]", source)
    right = evaluate_grounding("ISMS 인증을 준비하며 2026년 하반기 신청 예정이다. [SEC]", source)
    assert wrong.soft_flags and right.decision == "ACCEPT" and not right.soft_flags
    from esgenie.supplychain.drafter import _attempt_draft
    from esgenie.supplychain.schema import Answer
    # 부분 관련 근거는 확인된 신원보호 문장까지만 초안으로 사용한다.
    sentence = "신고자의 신원은 엄격히 보호하며 익명 신고를 허용한다."
    assert sentence in texts[CLAUSE_DOCS[2]]
    chunks, _ = select_fit_chunks([{"id":"DOC17", "text":texts[CLAUSE_DOCS[2]], "source_file":CLAUSE_DOCS[2], "page":0}], fmap["HMC-E-8"])
    fw, _ = hmc_fitness()
    q = next(q for q in fw.questions if q.qid == "HMC-E-8")
    a = Answer(q.qid,q.section,q.text,None,"insufficient")
    llm = MagicMock()
    llm.complete.return_value = SimpleNamespace(content=sentence+" [DOC17]")
    _attempt_draft(a,chunks,llm,max_retries=0,fitness=fmap[q.qid])
    assert a.status == "draft_ready" and a.draft_grounding["decision"] == "ACCEPT"
    return {"input_class":"real PDF local text; manual relevance labels; actual grounding; no live LLM", "documents":rows, "mixed":mixed, "limited_e8_draft":a.to_dict(), "grounding":{"wrong_acquisition":wrong.to_dict(),"correct_preparation":right.to_dict()}}


def render_pdf(path, out_dir):
    import fitz
    shots=[]
    with fitz.open(path) as doc:
        for stale in out_dir.glob("pdf_p[0-9]*.png"):
            if int(stale.stem[5:]) > len(doc):
                stale.unlink()
        for i,page in enumerate(doc):
            dest=out_dir/f"pdf_p{i+1}.png"
            page.get_pixmap(dpi=110).save(dest)
            shots.append(str(dest))
    return shots
def survey_all_frameworks() -> dict:
    from esgenie.knowledge.kesg_evidence_requirements import requirement_for_question
    from esgenie.supplychain.frameworks import all_framework_keys, get_framework
    from esgenie.supplychain.question_fitness import build_fitness_map

    result = {}
    for key in all_framework_keys():
        fw = get_framework(key)
        fmap = build_fitness_map(
            fw.questions,
            lambda q: requirement_for_question(q.kesg_codes, quantitative=q.qtype == "numeric"),
        )
        counts = sorted(len(f.distinctive) for f in fmap.values())
        unusable = [qid for qid, f in fmap.items() if not f.usable]
        shared = {}
        for q in fw.questions:
            shared.setdefault(q.primary_code, []).append(q.qid)
        worst = max(shared.items(), key=lambda kv: len(kv[1]))
        result[key] = {
            "label": fw.label,
            "questions": len(fw.questions),
            "distinctive_min": counts[0],
            "distinctive_median": counts[len(counts) // 2],
            "distinctive_max": counts[-1],
            "questions_without_distinctive_terms": unusable,
            "most_shared_primary_code": {"code": worst[0], "question_count": len(worst[1])},
        }
    return result



def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--evidence-dir", type=Path, required=True)
    parser.add_argument("--cache-dir", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, default=OUT_DIR)
    parser.add_argument("--log-dir", type=Path, default=LOG_DIR)
    args=parser.parse_args()
    inputs, provenance=load_actual_inputs(args.evidence_dir,args.cache_dir)
    from esgenie.supplychain.claims import SupplierClaim
    saq_path=args.evidence_dir/"05_OEM_ESG자가진단설문_한성모터스.pdf"
    saq=pdf_text(saq_path)
    assert re.search(r"재활용률\s*92\s*%",saq)
    original_claim=SupplierClaim(code="E-6-2",value=92,unit="%",raw="재활용률 92% 달성",source="saq:"+saq_path.name)
    graph,extraction,points,sheet,evaluations=pipeline(inputs,{"E-6-2":original_claim})
    sheet=draft_actual_communication(sheet,graph)
    original=export_and_reopen(sheet,args.out_dir/"actual",controlled=False)
    waste=graph.resolved_facts["E-6-2"]
    claim=SupplierClaim(code="E-6-2",value=92,unit="%",raw="동일 4월·제1공장 범위 재활용률 92% (통제 주장)",source="synthetic:same_scope_control; original SAQ value only",boundary=waste.boundary)
    cg,ce,cp,cs,ev=pipeline(inputs,{"E-6-2":claim},label="동일범위주장통제")
    cs=draft_actual_communication(cs,cg)
    controlled=export_and_reopen(cs,args.out_dir/"controlled",controlled=True)
    dump(args.out_dir/"pipeline.json",{"inputs":[i.to_dict() for i in inputs],"graph":graph.to_dict(),"d1":evaluations,"data_points":[p.to_dict() for p in points],"controlled_data_points":[p.to_dict() for p in cp]})
    report={"import_path":str(Path(sys.modules['esgenie'].__file__).resolve()),"inputs":provenance,
        "claim_source":{"path":str(saq_path),"sha256":sha256(saq_path),"page":0,"actual_scope":"SAQ scope not proven identical to April waste ledger", "control":"value 92 reused, April/site boundary explicitly assigned only in synthetic claim"},
        "actual":original,"controlled":controlled,"negative_controls":negative_controls(sheet_payload(cs)),
        "clauses":replay_real_documents(args.evidence_dir),"framework_structure_only":survey_all_frameworks(),
        "pdf_renders":render_pdf(controlled['files']['pdf']['path'],args.out_dir/"controlled"),
        "generation":{"LLM":"deterministic exact-source excerpt mock","grounding":"production local gate, no mock","live_calls":0},
    }
    report['failed']=original['failed']+controlled['failed']
    dump(args.log_dir/"validation_report.json",report)
    print(json.dumps({"failed":report['failed'],"actual_checks":original['executed'],"controlled_checks":controlled['executed'],"negative_controls":report['negative_controls'],"out_dir":str(args.out_dir)},ensure_ascii=False,indent=2))
    if report['failed']: raise SystemExit(1)

if __name__ == "__main__":
    main()
