"""BM 개편 세트 실행 결과(그래프·응답서·LLM 본문·Excel·데이터시트·PDF)를 원문 정답과 항목별로 대조한다(제품 코드 비의존).

사용:
  python check_outputs.py <run>/<stage> [--variant] [--json out.json]      # 정상 출력 검사 — 하나라도 실패하면 종료 코드 1
  python check_outputs.py <run>/<stage> [--variant] --inject [--json …]    # 오답 주입 사본마다 검사기가 실패하는지 확인

PR71 후속(2026-10-05) 보강. 이전 판(4972df2)은 특정 오답 문자열과 'MD·PDF에 표시가 하나라도 있는가'만 보아, 잘못된
15명·21명 문장의 표시를 지워도 30/30이었다(검토 `probe_output_checker.py`). 이번 판은
  - 정답(원본 PDF 직접 판독, `00_baseline/expected_answers.md`)의 값·역할·날짜·범위·상태를 항목마다 고정하고,
  - 본문·PDF의 **모든 교육 인원 수량**을 그 정답과 대조한다 — 경고(`[검토: …]`)가 붙어도 오답은 오답이다,
  - 미확정·부분값 상태는 **그 값을 말한 같은 문장·제목·표 행**에 있어야 한다(다른 문장의 표시로 통과하지 않는다),
  - 오답 주입 사본(틀린 수치·날짜·범위, 상태 표시 제거, 잘린 사유)마다 해당 항목이 실패하는지 확인한다.
원본 산출물은 읽기만 한다. 주입은 메모리 사본(읽어 들인 텍스트·표·JSON)에만 한다.
검사 항목의 분모는 '고정한 정답 항목 수'이다 — 본문 전체 문장의 정확도가 아니다(본문은 교육 인원·범위·상태 항목만 본다).

PR71 재검토(29defa0) D 보강: 이전 판은 내역 수치를 숫자 집합 `{40, 8, 2, 6}`으로만 허용해 `정규직 6명과 기간제 40명이
출석`을 통과시켰다(검토 `probe_output_checker.py` 38/38). 이제 원본 출석대장을 직접 센 **값·고용형태·역할·날짜** 정답표
(`TRAINING_GROUPS`, `derive_training_groups.py`)와 대조하고, 집단을 특정할 수 없는 내역 수치는 실패로 본다. 다른 사업장
표기에 붙은 정답 수치, SOURCE_ONLY 항목의 사유 누락, 응답서·Excel·응답서 PDF의 교육 인원도 대조한다. 오답은 Markdown·PDF
양쪽(서로 일치)·한쪽에만 넣어 각각 해당 산출물 항목이 실패하는지 확인한다(파일끼리 비교하지 않는다).
"""
from __future__ import annotations

import argparse
import copy
import glob
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import fitz
import openpyxl

KEY = {
    "electricity_kwh": 142560, "gas_m3": 8420, "water_m3": 680,
    "waste_total_kg": 18400, "waste_recycled_kg": 5400, "waste_rate_pct": 29.3,
    "scrap_generated_kg": 12500, "scrap_reused_kg": 11500, "scrap_exported_kg": 1000, "scrap_reuse_pct": 92.0,
}
# 09·13 교육 기록(원본 1~3쪽 집계·대장, 13 대조표 직접 판독). 역할·날짜별 정답 인원.
TRAINING = {
    ("대상", "2026-04-22"): 50, ("참석", "2026-04-22"): 46, ("미참석", "2026-04-22"): 4,
    ("추가 참석", "2026-04-27"): 4, ("중복 제외 합계", "2026-04"): 50,
}
# 같은 기록의 고용형태별 내역 — 09 1쪽 `대상 50명: 정규직 40 + 기간제 8 + 파견 2`와 09·13 출석대장(인원번호별 구분·출석
# 상태)을 직접 센 값(`derive_training_groups.py` → `06_second_followup/training_groups_from_pdf.json`과 같다).
# 4/22 참석 = 정규직 40(HN-G01~40) + 기간제 6(HN-G41~46), 미참석 = 기간제 2 + 파견 2, 4/27 추가 = 기간제 2 + 파견 2.
TRAINING_GROUPS = {
    ("대상", "정규직", "2026-04-22"): 40, ("대상", "기간제", "2026-04-22"): 8, ("대상", "파견", "2026-04-22"): 2,
    ("참석", "정규직", "2026-04-22"): 40, ("참석", "기간제", "2026-04-22"): 6,
    ("미참석", "기간제", "2026-04-22"): 2, ("미참석", "파견", "2026-04-22"): 2,
    ("추가 참석", "기간제", "2026-04-27"): 2, ("추가 참석", "파견", "2026-04-27"): 2,
}
GROUP_WORD = r"(비정규직|정규직|기간제|계약직|파견)"
GROUP_NORMAL = {"계약직": "기간제"}
# 원문 사업장(정답 수치는 모두 김해 제1공장). 다른 사업장 표기에 붙은 정답 수치는 그 사업장 실적이라는 오답이다.
SITE_RE = re.compile(r"(?:(?P<region>[가-힣]{2,4})\s*)?제\s*(?P<number>\d+)\s*공장|(?P<name>양산|부산|김해)\s*공장")
# SOURCE_ONLY 사유(무엇을 확인하지 못했는지) 문구 — 상태 낱말(`범위 확인 필요`)만 남고 사유가 빠지면 실패다.
REASON_RE = re.compile(r"원문(?:에|의)[^|]{4,90}?(?:확인하지\s*못했습니다|확정하지\s*못했습니다|확인할\s*수\s*없)")
STATUS_WORDS = (r"범위\s*미확정|미확정|참고값|원문\s*범위로?만|확인하지\s*못|확인되지\s*않|확정하지\s*않|아닙니다|아님|아니다"
                r"|명확하지\s*않|불명확|확정\s*여부")
TRAINING_CONTEXT = r"교육|참석|출석|미참석|불참|이수|수료"
HOLD_MARK = "[확인 보류]"
# 파일명 안의 날짜(`09_…_2026-04-22.pdf`)는 서술의 날짜가 아니다.
FILE_NAME_RE = re.compile(r"\d{2}_[^\s|;,()]+?\.(?:pdf|xlsx)")


def glyph_flat(text: str) -> str:
    """PDF 내보내기는 글꼴에 없는 ÷·×·→를 `/`·`x`·`->`로 그린다 — 같은 글자로 맞춘 뒤 공백을 지운다."""
    text = str(text or "").replace("÷", "/").replace("×", "x").replace("→", "->")
    return re.sub(r"\s+", "", text)


@dataclass
class Bundle:
    """검사 대상 산출물을 읽어 둔 사본. 주입 검사는 이 사본만 바꾼다."""
    stage: Path
    pipeline: dict
    sheet: dict
    body: str = ""
    report_pdf_text: str = ""
    report_pdf_pages: int = 0
    sheet_pdf_text: str = ""
    sheet_pdf_pages: int = 0
    excel_rows: dict = field(default_factory=dict)
    datasheet: dict = field(default_factory=dict)
    # 추가 양식(kesg61) 응답서 — 변형본 E-8-1처럼 기본 양식(rba42)에 없는 문항.
    sheet_kesg61: dict | None = None
    excel_rows_kesg61: dict = field(default_factory=dict)
    body_review: list | None = None
    paths: dict = field(default_factory=dict)


def _first(pattern: str) -> str | None:
    return next(iter(sorted(glob.glob(pattern))), None)


def pdf_text(path):
    if not path:
        return "", 0
    doc = fitz.open(path)
    return "\n".join(page.get_text() for page in doc), doc.page_count


def load(stage: Path) -> Bundle:
    exports = stage / "exports"
    paths = {"md": _first(str(exports / "*" / "ESG보고서_*.md")),
             "report_pdf": _first(str(exports / "*" / "ESG보고서_*.pdf")),
             "datasheet": _first(str(exports / "*" / "ESG_DataSheet_*.xlsx")),
             "sheet_xlsx": _first(str(exports / "response_sheet" / "*.xlsx")),
             "sheet_pdf": _first(str(exports / "response_sheet" / "*.pdf")),
             "body_review": _first(str(exports / "*" / "report_body_review.json"))}
    b = Bundle(stage, json.loads((stage / "pipeline.json").read_text(encoding="utf-8")),
               json.loads((stage / "result.json").read_text(encoding="utf-8"))["sheet"], paths=paths)
    b.body = Path(paths["md"]).read_text(encoding="utf-8") if paths["md"] else ""
    b.report_pdf_text, b.report_pdf_pages = pdf_text(paths["report_pdf"])
    b.sheet_pdf_text, b.sheet_pdf_pages = pdf_text(paths["sheet_pdf"])
    if paths["sheet_xlsx"]:
        ws = openpyxl.load_workbook(paths["sheet_xlsx"])["응답서"]
        for r in ws.iter_rows(min_row=5, values_only=True):
            if r and r[0]:
                b.excel_rows[str(r[0])] = [str(c) if c is not None else "" for c in r]
    if paths["datasheet"]:
        ws = openpyxl.load_workbook(paths["datasheet"])["DataSheet"]
        head = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            b.datasheet[str(r[0])] = dict(zip(head, r))
    if paths["body_review"]:
        b.body_review = json.loads(Path(paths["body_review"]).read_text(encoding="utf-8"))
    if (stage / "result_kesg61.json").exists():
        b.sheet_kesg61 = json.loads((stage / "result_kesg61.json").read_text(encoding="utf-8"))["sheet"]
    x61 = _first(str(exports / "response_sheet_kesg61" / "*.xlsx"))
    if x61:
        ws = openpyxl.load_workbook(x61)["응답서"]
        for r in ws.iter_rows(min_row=5, values_only=True):
            if r and r[0]:
                b.excel_rows_kesg61[str(r[0])] = [str(c) if c is not None else "" for c in r]
    return b


def nodes_of(pipeline, prefix):
    return [n for n in pipeline["evidence_graph"]["nodes"] if (n.get("source_file") or "").startswith(prefix)]


def zero_status(node):
    for p in (node.get("boundary") or {}).get("provenance", []):
        if isinstance(p, dict) and p.get("source") == "zero_evidence":
            return f"{p.get('status')}/{p.get('cause')}"
    return ""


# ── 본문 문장 단위 판독(제품 코드와 독립) ─────────────────────────────────────

def narrative_units(text: str) -> list[str]:
    """보고서 본문을 판정 단위(문장·제목·표 행)로 나눈다. 시스템 원장 표(`| K-ESG |` 머리)와 인용 상자(`>`)는 뺀다."""
    units, system_table = [], False
    for line in text.splitlines():
        # 확인 필요 사항은 Markdown 특수 문자를 `\(`처럼 가린다 — 글자 그대로 읽는다(PDF에는 가림 표시가 없다).
        s = re.sub(r"\\([\\`*_{}\[\]()#+.!|<>-])", r"\1", line.strip())
        if not s:
            system_table = False
            continue
        if s.startswith("|"):
            if re.match(r"\|\s*K-ESG\s*\|", s):
                system_table = True
            if not system_table and not re.fullmatch(r"\|?[\s:|-]+\|?", s):
                units.append(s)
            continue
        system_table = False
        if s.startswith(">"):
            continue
        if s.startswith("#"):
            units.append(s)
            continue
        units += _attach_marks([p.strip() for p in re.split(r"(?<=[.!?。])\s+(?=\S)", s) if p.strip()])
    return units


def _attach_marks(parts: list[str]) -> list[str]:
    """문장 뒤에 붙은 `[검토: …]` 표시는 그 문장과 같은 자리다 — 마침표에서 갈렸어도 앞 문장에 붙인다."""
    out: list[str] = []
    for part in parts:
        if out and part.startswith("[검토:"):
            out[-1] += " " + part
        else:
            out.append(part)
    return out


def pdf_units(text: str) -> list[str]:
    flat = re.sub(r"\s+", " ", text.replace("\n", " "))
    return _attach_marks([p.strip() for p in re.split(r"(?<=[.!?。])\s+(?=\S)", flat) if p.strip()])


def is_source_quote(unit: str) -> bool:
    """부록(확인 필요 사항)의 원문 인용 줄 — `출처: 파일 · N쪽 …`. 생성 서술이 아니라 원문을 옮긴 것이다."""
    return bool(re.search(r"출처:\s*\S+\s*·\s*\d+쪽", unit))


def _dates(unit: str) -> set[str]:
    out = set()
    for m in re.finditer(r"(?:(20\d{2})\s*[-.년]\s*)?(\d{1,2})\s*[-.월]\s*(\d{1,2})\s*일?", unit):
        if m.group(2) and m.group(3) and 1 <= int(m.group(2)) <= 12 and 1 <= int(m.group(3)) <= 31:
            out.add(f"04-{int(m.group(3)):02d}" if int(m.group(2)) == 4 else f"{int(m.group(2)):02d}-{int(m.group(3)):02d}")
    return out


def _held_numbers_span(unit: str) -> tuple[int, int] | None:
    """확인 보류 문구가 인용한 모델 수치(`수치(15명, 6명)`) 구간 — 단정이 아니다."""
    m = re.search(r"생성 문장의 수치\(([^)]*)\)", unit)
    return m.span() if m else None


def _group_of(unit: str, start: int, end: int) -> tuple[str, bool]:
    """(집단, 바로 꾸밈 여부). 바로 앞 집단 낱말(`정규직 40명`·`기간제 근로자는 6명`) 또는 뒤 서술어(`40명은 정규직이다`).
    없으면 같은 절 앞에 집단 낱말이 하나뿐일 때 그것(바로 꾸미지 않음), 여럿·없음이면 ""."""
    before = unit[max(0, start - 16):start]
    m = re.search(GROUP_WORD + r"\s*(?:직|근로자|직원|인원|근무자)?\s*(?:수)?\s*(?:은|는|이|가|의|:|：)?\s*$", before)
    if m:
        return GROUP_NORMAL.get(m.group(1), m.group(1)), True
    m = re.match(r"\s*(?:은|는|이|가)\s*" + GROUP_WORD + r"(?:직|근로자|직원)?\s*(?:이다|이며|이고|이었|였|입니다|임)",
                 unit[end:end + 20])
    if m:
        return GROUP_NORMAL.get(m.group(1), m.group(1)), True
    clause = re.split(r"[,;.]|\d+\s*명", unit[max(0, start - 40):start])[-1]
    found = {GROUP_NORMAL.get(g, g) for g in re.findall(GROUP_WORD, clause)}
    return (found.pop(), False) if len(found) == 1 else ("", False)


def _role_of(context: str, unit: str, dates: set[str]) -> str:
    if re.search(r"미\s*참석|불참|결석|참석하지|출석하지", context):
        return "미참석"
    if re.search(r"추가", context) or "04-27" in dates:
        return "추가 참석" if re.search(r"참석|출석|참여|이수", context + unit) else ""
    if re.search(r"대상", context):
        return "대상"
    return "참석" if re.search(r"(?<!미)(?<!미 )(참석|출석|참여)", context) else ""


def training_count_errors(unit: str) -> list[str]:
    """한 판정 단위 안의 교육 인원 수량을 정답(역할·날짜, 고용형태별 내역은 값·집단·역할·날짜)과 대조한 오류 목록."""
    unit = FILE_NAME_RE.sub(lambda m: " " * len(m.group()), unit)
    if not re.search(TRAINING_CONTEXT, unit):
        return []
    skip = _held_numbers_span(unit)
    errors = []
    totals = set(TRAINING.values())
    for m in re.finditer(r"(?<![\d.,A-Za-z가-힣\-])(\d{1,3})\s*명", unit):
        if skip and skip[0] <= m.start() < skip[1]:
            continue
        # 교육 문맥은 그 수량 가까이(앞 30자·뒤 20자)에서만 본다 — PDF에서 다른 표(`이사 출석률`)와 붙은 다른 지표의
        # 인원(`직접 고용 48명`)을 교육 인원으로 읽지 않는다.
        if not re.search(TRAINING_CONTEXT, unit[max(0, m.start() - 30):m.end() + 20]):
            continue
        n = int(m.group(1))
        # 날짜는 그 수량 앞쪽 같은 서술(40자 안)에서만 읽는다 — 긴 표·부록 단위의 다른 날짜를 끌어오지 않는다.
        dates = _dates(unit[max(0, m.start() - 40):m.start()])
        # 역할 낱말은 같은 절에서만 본다 — 쉼표·다음 수량 너머(`46명이며, 미참석 인원은 4명`)의 말은 다른 수량의 것이다.
        before = re.split(r"[,;]|\d+\s*명", unit[max(0, m.start() - 14):m.start()])[-1]
        after = re.split(r"[,;]|\d+\s*명", unit[m.end():m.end() + 14])[0]
        context = before + "|" + after
        group, direct = _group_of(unit, m.start(), m.end())
        if group and (direct or n not in totals):
            # 고용형태별 내역: 값·집단이 정답표에 있고, 읽힌 역할·날짜와도 맞아야 한다(숫자 집합으로 허용하지 않는다).
            # 날짜는 그 수량에 가장 가까운 앞 날짜 하나다(앞 서술의 다른 날짜를 섞지 않는다).
            near = [d for d in re.finditer(r"(?:(20\d{2})\s*[-.년]\s*)?(\d{1,2})\s*[-.월]\s*(\d{1,2})\s*일?",
                                           unit[max(0, m.start() - 40):m.start()])]
            dates = _dates(near[-1].group(0)) if near else set()
            role = _role_of(context, unit, dates)
            fits = [key for key, count in TRAINING_GROUPS.items() if count == n and key[1] == group
                    and (not role or key[0] == role) and (not dates or key[2][5:] in dates)]
            if not fits:
                errors.append(f"고용형태별 인원 오기 {group} {n}명(역할 {role or '미상'}·날짜 {sorted(dates) or '미상'}): "
                              f"{unit[:120]}")
            continue
        if n not in totals:
            parts = {count for count in TRAINING_GROUPS.values()}
            errors.append(("집단을 특정할 수 없는 내역 수치" if n in parts else "정답에 없는 인원") + f" {n}명: {unit[:120]}")
            continue
        absent = re.search(r"미\s*참석|불참|결석", context)
        attended = re.search(r"(?<!미)(?<!미 )(참석|출석)", context) and not absent
        if n == 46 and absent:
            errors.append(f"참석 46명을 미참석으로 서술: {unit[:120]}")
        if n == 4 and attended and not re.search(r"추가|27일|04-27", unit) and not absent:
            errors.append(f"미참석 4명을 참석으로 서술: {unit[:120]}")
        if n == 50 and attended and not re.search(r"대상|중복|고유|합계|추가|27일|04-27|전원", context + unit) \
                and ("04-22" in dates or not dates):
            errors.append(f"대상 50명을 4/22 참석으로 서술: {unit[:120]}")
        if n == 50 and "04-22" in dates and re.search(r"중복\s*제외|고유\s*인원", context) and "04-27" not in dates \
                and not re.search(r"원문 확인 값", unit):
            errors.append(f"4/27 이후 중복 제외 50명을 4/22 값으로 서술: {unit[:120]}")
        if n == 46 and "04-27" in dates and "04-22" not in dates:
            errors.append(f"4/22 참석 46명을 4/27 값으로 서술: {unit[:120]}")
    if re.search(r"ID\s*50\s*개|50\s*개[^.]{0,6}ID", unit) and re.search(r"검증|확인", unit):
        errors.append(f"명단 ID 50개 검증 서술: {unit[:120]}")
    # 부정(`판정하는 문서는 아닙니다`·`충족했다는 판정은 아닙니다`)은 확대가 아니다.
    if re.search(r"법정[^.]{0,12}(?:교육|이수)[^.]{0,12}(?:완료|충족|이수했)", unit) and not re.search(r"아니|아닙|않|아님", unit):
        errors.append(f"법정 교육 전체 이수로 확대: {unit[:120]}")
    return errors


KEY_VALUE_RE = re.compile(r"(?<![\d.,])(?:142,?560|0\.513216|68\.158|680(?:\.0)?\s*(?:톤|m³|ton)|18,?400|18\.4\s*톤|29\.3\s*%"
                          r"|(?:50|46|4)\s*명)")


def site_errors(units: list[str]) -> list[str]:
    """정답 수치(교육 인원·4월 제1공장 환경값)를 다른 사업장 표기에 붙여 쓴 단위. 수량 앞 25자 안의 가장 가까운 사업장 표기를 본다."""
    bad = []
    for unit in units:
        if is_source_quote(unit):
            continue
        for m in KEY_VALUE_RE.finditer(unit):
            if m.group(0).endswith("명") and not re.search(TRAINING_CONTEXT, unit[max(0, m.start() - 30):m.end() + 20]):
                continue
            sites = list(SITE_RE.finditer(unit[max(0, m.start() - 25):m.start()]))
            if not sites:
                continue
            site = sites[-1]
            # 원문의 다른 사업장은 양산 제2공장이다. 앞 낱말(`같은 기간 제1공장`)을 지역명으로 읽지 않는다 — 번호·알려진 지역만 본다.
            if (site.group("number") and site.group("number") != "1") or site.group("name") in ("양산", "부산") \
                    or site.group("region") in ("양산", "부산"):
                bad.append(f"다른 사업장 표기({site.group(0)})에 정답 수치 {m.group(0)}: {unit[:120]}")
    return bad


GROUP_ERRORS = ("고용형태별 인원 오기", "집단을 특정할 수 없는 내역 수치")


def sheet_text_units(sheet: dict, excel_rows: dict, sheet_pdf_text: str) -> dict[str, list[str]]:
    """응답서 JSON(근거·사유·표시)·Excel 행·응답서 PDF를 문장 단위로 나눈다 — 교육 인원 대조용."""
    def split(text: str) -> list[str]:
        return pdf_units(text)
    answer = []
    for a in sheet.get("answers", []):
        answer += split(" ".join(str(a.get(k) or "") for k in ("rationale", "comparison_reason", "review_note"))
                        + " " + " ".join(a.get("flags") or []))
    return {"answer": answer, "excel": [u for row in excel_rows.values() for u in split(" | ".join(row))],
            "sheet_pdf": split(sheet_pdf_text)}


def fact_listing_errors(unit: str) -> list[str]:
    """확인 보류 문구의 '원문 확인 값' 목록(`라벨 N명(날짜)`)도 정답과 대조한다."""
    m = re.search(r"원문 확인 값:\s*(.*)", unit)
    if not m:
        return []
    errors = []
    for item in re.split(r";\s*", m.group(1)):
        item = item.split(" — ")[0]          # 뒤의 출처 파일명은 역할 판독에 쓰지 않는다
        q = re.search(r"(\d{1,3})명\((20\d{2}-\d{2}(?:-\d{2})?)", item)
        if not q:
            continue
        n, when = int(q.group(1)), q.group(2)
        if "미참석" in item:
            role = "미참석"
        elif "추가" in item:
            role = "추가 참석"
        elif re.search(r"중복|고유 인원\s*50|합계", item) and "22일" not in item:
            role = "중복 제외 합계"
        elif "대상" in item:
            role = "대상"
        elif re.search(r"참석|출석", item):
            role = "참석"
        else:
            continue
        expected = TRAINING.get((role, when))
        if expected is None or expected != n:
            errors.append(f"원문 확인 값 오기: {item.strip()[:100]} (정답 {role}·{when} = {expected})")
    return errors


def status_missing(units: list[str], mention: str, value_re: str) -> list[str]:
    """`mention`의 값을 말한 단위마다 같은 단위에 미확정 상태가 있는가(부정된 유보는 상태가 아니다)."""
    missing = []
    for unit in units:
        if not (re.search(mention, unit) and re.search(value_re, unit)):
            continue
        hedge = [h for h in re.finditer(STATUS_WORDS, unit)
                 if not re.match(r"\s*(?:할|될|한|된)?\s*(?:사항|것|부분)?\s*(?:이|은|는)?\s*없", unit[h.end():])]
        if not hedge or re.search(r"(?<![미불])확정(?:되었|됐|된\s*실적|됨)", unit):
            missing.append(unit[:160])
    return missing


def widened_scope(units: list[str], value_re: str) -> list[str]:
    bad = []
    for unit in units:
        if re.search(value_re, unit) and re.search(r"사업장\s*전체|전사|연간|합산|모든\s*사업장", unit) \
                and not re.search(r"부분|월간|아닙니다|아님|아니|불명확|한정|원장 범위", unit):
            bad.append(unit[:160])
    return bad


# ── 항목 검사 ────────────────────────────────────────────────────────────────

def check(stage: Path | Bundle, *, variant: bool = False) -> dict:
    b = stage if isinstance(stage, Bundle) else load(stage)
    pipeline, sheet = b.pipeline, b.sheet
    followup = b.stage.name == "followup" or b.stage.name.endswith("followup")
    rows = []

    def add(item, ok, observed, expected, where):
        rows.append({"item": item, "pass": bool(ok), "observed": observed, "expected": expected, "where": where})

    # ── 그래프·추출 ─────────────────────────────────────────────
    e02 = [n for n in nodes_of(pipeline, "02_") if n["metric"] == "E-4-1"]
    add("02 전기 142,560 kWh(E-4-1)", any(n["value"] == KEY["electricity_kwh"] and n["unit"] == "kWh" for n in e02),
        [(n["value"], n["unit"]) for n in e02], "142560 kWh", "graph")
    g03 = nodes_of(pipeline, "03_")
    add("03 가스 8,420 m³, 열량·배출량 미생성", any(n["value"] == KEY["gas_m3"] and n["unit"] in ("m³", "m3") for n in g03)
        and not any(n["unit"] in ("MJ", "GJ", "TJ") or n["metric"] in ("E-4-1", "E-3-1") for n in g03),
        [(n["metric"], n["value"], n["unit"]) for n in g03], "8420 m³ only", "graph")
    g07 = nodes_of(pipeline, "07_")
    add("07 수도 680 m³", any(n["value"] == KEY["water_m3"] for n in g07), [(n["metric"], n["value"], n["unit"]) for n in g07],
        "680 m³", "graph")
    g04 = nodes_of(pipeline, "04_")
    vals04 = {(n["metric"], n["value"], n["unit"]) for n in g04}
    add("04 총 위탁량 18,400 kg·재활용량 5,400 kg·재활용률 29.3%",
        any(v == KEY["waste_total_kg"] and u == "kg" for _, v, u in vals04)
        and any(v == KEY["waste_recycled_kg"] and u == "kg" for _, v, u in vals04)
        and any(v == KEY["waste_rate_pct"] and u == "%" for _, v, u in vals04), sorted(vals04), "3값", "graph")
    add("06 규정에서 수량 미생성(3톤 없음)", not nodes_of(pipeline, "06_"),
        [(n["metric"], n["value"], n["unit"]) for n in nodes_of(pipeline, "06_")], "노드 0", "graph")
    g08 = nodes_of(pipeline, "08_")
    rate08 = [n for n in g08 if n["unit"] == "%"]
    add("08 내부 재투입률 92%는 E-2-2·E-6-2가 아님", rate08 and all(n["metric"] not in ("E-2-2", "E-6-2") for n in rate08),
        [(n["metric"], n["value"]) for n in rate08], "코드 없음(원문 라벨)", "graph")
    totals = {n["value"]: n["metric"] for n in g08 if n["value"] in (12500, 11500, 1000)}
    add("08 발생 12,500·재투입 11,500·반출 1,000 라벨 구분",
        "발생" in totals.get(12500, "") and "재투입" in totals.get(11500, "") and "반출" in totals.get(1000, ""),
        totals, "발생/재투입/반출 라벨", "graph")
    g09 = nodes_of(pipeline, "09_")
    by_value = {}
    for n in g09:
        if n["unit"] == "명":
            by_value.setdefault(n["value"], []).append(n["metric"])
    add("09 50명 라벨 = 대상(참석·출석 아님)", by_value.get(50) and all("대상" in m and not re.search(r"(?<!미)(참석|출석)\s*인원$", m)
                                                                for m in by_value[50]), by_value.get(50), "대상", "graph")
    add("09 46명 라벨 = 참석·출석(미참석 아님)", by_value.get(46) and all(re.search(r"참석|출석", m) and "미참석" not in m
                                                               for m in by_value[46]), by_value.get(46), "참석", "graph")
    add("09 4명 라벨 = 미참석", by_value.get(4) and all("미참석" in m or "불참" in m for m in by_value[4]), by_value.get(4),
        "미참석", "graph")
    add("09 참석 비율 92%", any(n["value"] == 92 and n["unit"] == "%" for n in g09),
        sorted({(n["metric"], n["value"]) for n in g09 if n["unit"] == "%"}), "92%", "graph")
    g11 = nodes_of(pipeline, "11_")
    isms = [n for n in g11 if n["value"] == 0]
    clauses = [t for t in pipeline["evidence_graph"]["text_nodes"]
               if (t.get("source_file") or "").startswith("11_") and "ISMS" in t.get("text", "")
               and "미보유" in t.get("text", "") and "2026-04-30" in t.get("text", "")]
    unread = [f for f in pipeline.get("review_findings") or []
              if "인증 보유" in str(f.get("fact") or "") and "읽지 못한" in str(f.get("title") or "")]
    add("11 ISMS 인증 미보유·기준일 2026-04-30 보존(수치 0 또는 원문 조항)",
        (bool(isms) and all(zero_status(n).split("/")[0] in ("CONFIRMED", "SOURCE_ONLY") for n in isms) or bool(clauses))
        and not unread,
        {"zero_nodes": [(n["metric"], n["value"], zero_status(n)) for n in isms], "clauses": [t["text"][:60] for t in clauses],
         "unread_findings": len(unread)}, "0 또는 미보유 조항, 미확인 표시 없음", "graph")
    if followup:
        g13 = nodes_of(pipeline, "13_")
        m13 = {n["value"]: (n["metric"], ((n.get("boundary") or {}).get("period_text") or "")) for n in g13 if n["unit"] == "명"}
        add("13 추가 참석 4명 = 4월 27일 추가 참석", "추가" in m13.get(4, ("", ""))[0] and "27" in "".join(m13.get(4, ("", ""))),
            m13.get(4), "4월 27일 추가 참석", "graph")
        add("13 46명 = 4월 22일 참석(덮어쓰기 없음)", "22" in "".join(m13.get(46, ("", ""))) and "참석" in m13.get(46, ("", ""))[0],
            m13.get(46), "4월 22일 참석", "graph")
        add("13 50명 = 중복 제외 합계(대상자 합계 아님)", re.search(r"중복|고유", m13.get(50, ("", ""))[0]),
            m13.get(50), "중복 제외 합계", "graph")
    all_values = [n["value"] for n in pipeline["evidence_graph"]["nodes"]]
    add("이전 세트 360,772 MJ 미유입", 360772 not in all_values, "", "없음", "graph")
    if variant:
        e81 = [n for n in pipeline["evidence_graph"]["nodes"] if (n.get("source_file") or "").startswith("14_")]
        add("변형본 E-8-1 0건 = SOURCE_ONLY(회계연도 구간 미정의)",
            any(n["value"] == 0 and zero_status(n).startswith("SOURCE_ONLY") for n in e81),
            [(n["metric"], n["value"], zero_status(n)) for n in e81], "SOURCE_ONLY/fiscal_*", "graph")

    # ── 응답서 ─────────────────────────────────────────────────
    answers = {a["qid"]: a for a in sheet["answers"]}
    e62 = answers.get("RBA-C-4-E-6-2", {})
    reason = (e62.get("comparison_reason") or "") + " " + " ".join(e62.get("flags") or [])
    add("응답서 E-6-2 29.3% · 4월 제1공장 범위", e62.get("value") == 29.3 and "제1공장" in (e62.get("boundary_label") or ""),
        (e62.get("value"), e62.get("boundary_label")), "29.3 · 부분", "answer")
    add("응답서 05 회사 답변 92%를 일치·동일 범위 충돌로 확정하지 않음", e62.get("comparison") not in ("compared", "mismatch", None),
        (e62.get("status"), e62.get("comparison")), "not_comparable 또는 scope_unconfirmed", "answer")
    add("응답서 05 설명: 메모가 가리킨 08 내부 재투입률 계산식과 29.3% 계산식",
        all(t in reason for t in ("08_", "재투입", "11,500 kg ÷ 공정 스크랩 발생 12,500 kg", "5,400 kg ÷ 18,400 kg",
                                  "HW-SCR-202604")), reason[:300], "두 계산식 + 메모 ID", "answer")
    ctx = next(iter(e62.get("self_reports") or [{}])).get("context") or {}
    # 메모의 문서 ID가 원문에 있는 문서(이전 형식은 value_trace의 linked_by_note만 남겼다).
    linked = set((ctx.get("note_link") or {}).get("found_in") or []) | {
        t.get("source_file") for t in ctx.get("value_trace") or [] if t.get("linked_by_note")}
    named = set(re.findall(r"(\d{2}_[^\s']+\.pdf)", reason)) - {"04_사업장폐기물위탁처리명세_2026-04.pdf"}
    add("응답서 05 '옮긴 출처'는 메모가 직접 가리킨 문서뿐", named and named <= linked if "옮긴" in reason else True,
        {"named": sorted(named), "memo_linked": sorted(linked)}, "메모 연결 문서만", "answer")
    add("응답서 05 설명이 기간·사업장을 '미확인'으로 잘못 적지 않음", "기간·사업장·분모 미확인" not in reason, reason[:120],
        "기간·사업장 같음 명시", "answer")
    e41 = answers.get("RBA-C-8-E-4-1", {})
    add("응답서 E-4-1 0.513216 TJ 부분값(전력만)", e41.get("value") == 0.513216 and "부분" in (e41.get("boundary_label") or ""),
        (e41.get("value"), e41.get("boundary_label")), "0.513216 TJ · 부분", "answer")
    c2 = answers.get("RBA-C-2", {})
    add("응답서 C-2가 재생 원부자재 92%를 근거로 쓰지 않음", "재생 원부자재" not in (c2.get("rationale") or ""),
        (c2.get("status"), (c2.get("rationale") or "")[:80]), "E-2-2 92% 미사용", "answer")
    if variant:
        pool = (b.sheet_kesg61 or {}).get("answers", []) + sheet["answers"]
        e81a = next((a for a in pool if a["qid"].endswith("E-8-1")), {})
        add("응답서 변형본 E-8-1: 요청 범위 확정 실적이 아님(scope_unconfirmed, 원문 범위 표기)",
            e81a.get("status") != "verified" and e81a.get("comparison") == "scope_unconfirmed"
            and re.search(r"원문 범위로만|범위 확인 필요", " ".join(e81a.get("flags") or []) + (e81a.get("review_note") or "")),
            (e81a.get("qid"), e81a.get("status"), e81a.get("comparison")), "scope_unconfirmed", "answer")
        reason81 = " ".join(str(e81a.get(k) or "") for k in ("comparison_reason", "review_note"))
        add("응답서 변형본 E-8-1: 미확정 사유(무엇을 확인하지 못했는지) 보존", bool(REASON_RE.search(reason81)),
            reason81[:200], "원문 …을 확인하지 못했습니다", "answer")
    texts = sheet_text_units(sheet, b.excel_rows, b.sheet_pdf_text)
    for where, label in (("answer", "응답서"), ("excel", "Excel"), ("sheet_pdf", "응답서 PDF")):
        if where == "sheet_pdf" and b.sheet_pdf_pages == 0:
            continue
        wrong = [e for u in texts[where] for e in training_count_errors(u)]
        add(f"{label}의 교육 인원(원문 인용 포함)이 정답의 값·집단·역할·날짜와 일치", not wrong, wrong[:3],
            "정답 인원만", where)

    # ── LLM 본문(보고서 Markdown) ───────────────────────────────
    units = narrative_units(b.body)
    narrative = "\n".join(units)
    add("본문·요약에 '재생 원부자재 비율 92%' 없음", not re.search(r"재생\s*원부자재[^.\n]{0,20}92", narrative),
        re.findall(r"[^.\n]*재생\s*원부자재[^.\n]*", narrative)[:2], "없음", "report_md")
    errors = [e for u in units for e in training_count_errors(u) + fact_listing_errors(u)]
    group_errors = [e for e in errors if e.startswith(GROUP_ERRORS)]
    add("본문 교육 인원 수량이 모두 정답의 역할·날짜와 일치(경고가 붙은 오답도 실패)",
        not [e for e in errors if e not in group_errors], [e for e in errors if e not in group_errors][:6],
        "대상 50·참석 46·미참석 4(4/22), 추가 4(4/27), 중복 제외 50(4월)", "report_md")
    add("본문 교육 인원 고용형태별 내역이 정답표(값·집단·역할·날짜)와 일치", not group_errors, group_errors[:6],
        "정규직 40·기간제 6 참석, 기간제 2·파견 2 미참석(4/22) 등", "report_md")
    sites = site_errors(units)
    add("본문 정답 수치를 다른 사업장 실적으로 쓴 문장 없음", not sites, sites[:3], "김해 제1공장 수치만", "report_md")
    # 보존 여부는 사회 영역 본문(서술·표)에서만 본다 — 부록(확인 필요 사항)의 원문 인용으로 통과시키지 않는다.
    at = b.body.find("## 사회 성과")
    end = b.body.find("\n## ", at + 1) if at >= 0 else -1
    s_units = narrative_units(b.body[at:end if end > 0 else None]) if at >= 0 else []
    has46 = any(re.search(r"(?<!\d)46\s*명", u) and re.search(r"참석|출석", u) for u in s_units)
    add("사회 본문에 4/22 참석 46명 보존", has46, [u[:100] for u in s_units if "46명" in u][:3], "46명 서술", "report_md")
    if followup:
        has50 = any(re.search(r"(?<!\d)50\s*명", u) and re.search(r"중복|고유", u) for u in s_units)
        add("사회 본문에 추가 교육 후 중복 제외 50명 보존", has50, [u[:100] for u in s_units if "50명" in u][:3],
            "중복 제외 50명", "report_md")
    wide = widened_scope(units, r"0\.513216|0\.51\s*TJ|680(?:\.0)?\s*(?:톤|m³)|142,?560")
    add("본문 환경 수치(4월 제1공장 부분값)를 전체·연간·합산으로 넓힌 문장 없음", not wide, wide[:3], "없음", "report_md")
    claims_isms = [a["qid"] for a in sheet["answers"]
                   if re.search(r"ISMS[^.]{0,20}(?:보유|취득)(?!\s*하지|하지)", (a.get("rationale") or "") + " ".join(a.get("flags") or []))
                   and "미보유" not in (a.get("rationale") or "")]
    widened_isms = [u for u in units if re.search(r"ISMS[^.]{0,30}(?:취득|보유하고)", u) and "미보유" not in u]
    add("본문·응답서에 ISMS 인증 보유·취득 주장 없음", not claims_isms and not widened_isms,
        {"answers": claims_isms, "body": widened_isms[:2]}, "없음", "answer+report_md")
    if variant:
        miss = status_missing(units, r"환경\s*법규\s*위반|법규\s*위반|E-8-1", r"0\s*건|0\.0\s*건|제로|없었|없음|위반\s*없")
        add("본문 변형본 E-8-1 0건을 말한 문장·제목·표 행마다 같은 자리에 미확정 상태", not miss, miss[:3],
            "같은 단위에 미확정·참고값 상태", "report_md")
    held = [u for u in units if u.startswith(HOLD_MARK) or HOLD_MARK in u]
    if held or b.body_review is not None:
        replaced = [r for r in (b.body_review or []) if r.get("action") == "replaced"]
        add("본문에서 바꾼 문장마다 감사 기록에 모델 원문·사유 보존",
            b.body_review is not None and all(r.get("model_text") and r.get("reason") for r in replaced)
            and len({r.get("output") for r in replaced}) >= min(1, len(held)),
            {"held_units": len(held), "replaced_records": len(replaced)}, "report_body_review.json", "report_audit")

    # ── 응답서 Excel ──────────────────────────────────────────
    x62 = " | ".join(b.excel_rows.get("RBA-C-4-E-6-2", []))
    add("Excel E-6-2 행: 29.3%·범위·회사 답변 92%·출처", "29.3" in x62 and "제1공장" in x62 and "92" in x62 and "08_" in x62,
        x62[:400], "값·범위·회사 답변·출처", "excel")
    x41 = " | ".join(b.excel_rows.get("RBA-C-8-E-4-1", []))
    add("Excel E-4-1 행: 0.513216·부분값 범위", "0.513216" in x41 and "부분" in x41, x41[:200], "부분", "excel")
    if variant:
        x81 = " | ".join(next((v for k, v in {**b.excel_rows, **b.excel_rows_kesg61}.items() if k.endswith("E-8-1")), []))
        add("Excel 변형본 E-8-1 행: 같은 행에 원문 범위·확인 필요 상태", re.search(r"원문 범위로만|범위 확인 필요|미확정", x81)
            and "자가신고 일치" not in x81, x81[:300], "같은 행 상태", "excel")
        add("Excel 변형본 E-8-1 행: 같은 행에 미확정 사유", bool(REASON_RE.search(x81)), x81[:300], "같은 행 사유", "excel")
    ds = b.datasheet
    add("데이터시트 E-5-1·E-6-1 측정 범위 표시", all("제1공장" in str((ds.get(c) or {}).get("측정 범위") or "") for c in ("E-5-1", "E-6-1")),
        {c: (ds.get(c) or {}).get("측정 범위") for c in ("E-5-1", "E-6-1")}, "4월·제1공장", "datasheet")
    add("데이터시트 E-4-1 값·부분 상태", (ds.get("E-4-1") or {}).get("값") == 0.513216
        and "부분" in str((ds.get("E-4-1") or {}).get("범위·확정 상태") or ""),
        {k: (ds.get("E-4-1") or {}).get(k) for k in ("값", "범위·확정 상태")}, "0.513216 · 부분", "datasheet")
    if variant:
        # 데이터시트는 프로필(sme) 대상 항목만 싣는다 — E-8-1이 없으면 확정 실적으로 실린 것도 아니다. 있으면 상태가 있어야 한다.
        d81 = ds.get("E-8-1")
        add("데이터시트 변형본 E-8-1: 실리지 않았거나, 실렸으면 0과 원문 범위 보존 상태", d81 is None or (
            d81.get("값") == 0 and "원문 범위로만 보존" in str(d81.get("범위·확정 상태") or "")),
            {k: (d81 or {}).get(k) for k in ("값", "범위·확정 상태")} if d81 else "데이터시트 대상 아님(프로필 밖)",
            "없음 또는 0 · 원문 범위로만 보존", "datasheet")

    # ── PDF ───────────────────────────────────────────────────
    flat = glyph_flat(b.sheet_pdf_text)
    add("응답서 PDF 생성·E-6-2 29.3과 05 사유 전문(두 계산식·메모 ID·확정하지 않음)",
        b.sheet_pdf_pages > 0 and "29.3" in flat and all(glyph_flat(t) in flat for t in (
            "08_공정스크랩관리대장", "11,500 kg ÷ 공정 스크랩 발생 12,500 kg", "5,400 kg ÷ 18,400 kg", "HW-SCR-202604",
            "확정하지 않았습니다")),
        {"pages": b.sheet_pdf_pages}, "29.3·08 출처·두 계산식", "sheet_pdf")
    punits = pdf_units(b.report_pdf_text)
    rflat = glyph_flat(b.report_pdf_text)
    add("보고서 PDF에 '재생원부자재비율92' 없음", b.report_pdf_pages > 0 and not re.search(r"재생원부자재비율[^。]{0,8}92", rflat),
        {"pages": b.report_pdf_pages}, "없음", "report_pdf")
    perrors = [e for u in punits for e in training_count_errors(u) + fact_listing_errors(u)]
    pgroup = [e for e in perrors if e.startswith(GROUP_ERRORS)]
    add("보고서 PDF 교육 인원 수량이 정답과 일치", b.report_pdf_pages > 0 and not [e for e in perrors if e not in pgroup],
        [e for e in perrors if e not in pgroup][:4], "정답 인원만", "report_pdf")
    add("보고서 PDF 교육 인원 고용형태별 내역이 정답표와 일치", b.report_pdf_pages > 0 and not pgroup, pgroup[:4],
        "값·집단·역할·날짜", "report_pdf")
    psites = site_errors(punits)
    add("보고서 PDF 정답 수치를 다른 사업장 실적으로 쓴 문장 없음", b.report_pdf_pages > 0 and not psites, psites[:3],
        "김해 제1공장 수치만", "report_pdf")
    md_status = [u for u in units if re.match(r"\[(확인 보류|범위 미확정|범위 주의)\]", u) or "[검토:" in u]
    absent = []
    for u in md_status:
        probe = glyph_flat(re.sub(r"[#|*]", "", u))[:48]
        if probe and probe not in rflat:
            absent.append(u[:100])
    add("보고서 PDF에 본문의 보류·미확정 문구가 같은 내용으로 실림(문장별)", b.report_pdf_pages > 0 and not absent,
        {"md_status_units": len(md_status), "missing_in_pdf": absent[:3]}, "모두 실림", "report_pdf")
    if variant:
        pmiss = status_missing([u for u in punits if not is_source_quote(u)], r"환경\s*법규\s*위반|법규\s*위반",
                               r"0\s*건|0\.0\s*건|제로|없었|위반\s*없")
        add("보고서 PDF 변형본 E-8-1 0건 문장마다 같은 자리에 미확정 상태", b.report_pdf_pages > 0 and not pmiss, pmiss[:3],
            "같은 문장 상태", "report_pdf")
    groups = {}
    for r in rows:
        g = groups.setdefault(r["where"], [0, 0])
        g[0] += r["pass"]
        g[1] += 1
    return {"stage": str(b.stage), "passed": sum(r["pass"] for r in rows), "total": len(rows), "rows": rows,
            "by_artifact": {k: f"{v[0]}/{v[1]}" for k, v in groups.items()}}


# ── 오답 주입(메모리 사본) ──────────────────────────────────────────────────

def _sub_body(b: Bundle, pattern: str, repl: str, *, pdf: bool = True) -> int:
    """본문 **판정 단위(서술·제목·생성 표)** 안의 첫 일치만 바꾼다(부록 인용 상자·원장 표가 아니라). PDF도 같은 서술
    구간(`지표 해설` 뒤)의 첫 일치를 바꾼다."""
    n = 0
    # 교육 서술이 있는 사회 영역 본문을 먼저 본다(부록 `확인된 내용` 줄이 아니라).
    at = b.body.find("## 사회 성과")
    scope = b.body[at:] if at >= 0 else b.body
    for unit in narrative_units(scope):
        if re.search(pattern, unit):
            new = re.sub(pattern, repl, unit, count=1)
            if unit in b.body:
                b.body = b.body.replace(unit, new, 1)
                n = 1
                break
    if n == 0:
        b.body, n = re.subn(pattern, repl, b.body, count=1)
    if pdf:
        at = b.report_pdf_text.find("지표 해설")
        head, tail = (b.report_pdf_text[:at], b.report_pdf_text[at:]) if at >= 0 else ("", b.report_pdf_text)
        tail, _ = re.subn(pattern, repl, tail, count=1)
        b.report_pdf_text = head + tail
    return n


def injections(base: Bundle, *, variant: bool) -> list[tuple[str, str, callable]]:
    """(이름, 실패해야 하는 항목 접두, 사본 변경 함수). 각 함수는 바꾼 곳의 수를 돌려준다(0이면 주입 불가)."""
    cases = [
        ("본문 참석 46명 → 47명(틀린 수치)", "본문 교육 인원",
         lambda b: _sub_body(b, r"(?<!\d)46(\s*명)", r"47\1")),
        ("본문에 경고 붙은 오답 문장 추가(정규직 15명 출석 [검토: …])", "본문 교육 인원",
         lambda b: _sub_body(b, r"(### 주요 활동\n)", r"\12026년 4월 22일 교육에는 정규직 15명이 출석하였다. "
                             r"[검토: 근거 확인 필요: 인용 근거에서 찾지 못한 숫자 15 — 확정 사실로 보지 마세요] ", pdf=False)
         + _pdf_append(b, " 2026년 4월 22일 교육에는 정규직 15명이 출석하였다. [검토: 근거 확인 필요] ")),
        ("본문 참석·미참석 역할 뒤바꿈(46명이 미참석)", "본문 교육 인원",
         lambda b: _sub_body(b, r"(?<!\d)46(\s*명)", r"46\1이 미참석했으며")),
        ("본문 날짜 오기(4/22 참석 46명 → 4/27)", "본문 교육 인원",
         lambda b: _sub_body(b, r"(### 주요 활동\n)", r"\12026년 4월 27일 교육에는 46명이 참석하였다. ", pdf=False)),
        ("본문 환경 부분값을 연간 전사 값으로 확대", "본문 환경 수치",
         lambda b: _sub_body(b, r"(### 주요 활동\n)", r"\12026년 연간 전사 에너지 사용량은 0.513216 TJ이다. ", pdf=False)),
        ("응답서 E-6-2 값 29.3 → 92.0", "응답서 E-6-2 29.3%",
         lambda b: _set_answer(b, "RBA-C-4-E-6-2", "value", 92.0)),
        ("응답서 05 비교를 동일 범위 충돌(mismatch)로 바꿈", "응답서 05 회사 답변",
         lambda b: _set_answer(b, "RBA-C-4-E-6-2", "comparison", "mismatch")),
        ("Excel E-4-1 행에서 '부분' 범위 제거", "Excel E-4-1",
         lambda b: _excel_strip(b, "RBA-C-8-E-4-1", "부분")),
        ("응답서 PDF 05 사유 잘림(29.3% 계산식 누락)", "응답서 PDF",
         lambda b: _sheet_pdf_cut(b, "5,400 kg ÷ 18,400 kg")),
        ("보고서 PDF 참석 46명 → 45명", "보고서 PDF 교육 인원",
         lambda b: _pdf_sub(b, r"(?<!\d)46(\s*명)", r"45\1")),
        ("보고서 PDF에서 보류 문구 제거", "보고서 PDF에 본문의 보류",
         lambda b: _pdf_sub(b, r"\[(확인 보류|범위 미확정|범위 주의)\][^.]*\.", "", count=0)),
    ]
    if variant:
        cases += [
            ("변형본 본문 미확정 문장의 상태 문구 제거(값 0건 단정만 남김)", "본문 변형본 E-8-1",
             lambda b: _sub_body(b, r"\[범위 미확정\][^\n]*?(환경 법규 위반[^\n]*?0(?:\.0)?건)[^\n]*?\.",
                                 r"\1을 기록했습니다.", pdf=False)
             or _sub_body(b, r"(### 주요 활동\n)", r"\1FY2026 환경 법규 위반은 0건으로 확정되었습니다. ", pdf=False)),
            ("변형본 본문 '[검토: 범위 미확정 …]' 표시만 제거(문장은 그대로)", "본문 변형본 E-8-1",
             lambda b: _strip_marks(b, r"\s*\[검토: 범위 미확정[^\]]*\]")),
            ("변형본 데이터시트 E-8-1 상태 제거", "데이터시트 변형본 E-8-1",
             lambda b: _datasheet_strip(b, "E-8-1", "범위·확정 상태")),
            ("변형본 Excel E-8-1 행 상태 제거", "Excel 변형본 E-8-1",
             lambda b: _excel_strip(b, next((k for k in b.excel_rows_kesg61 if k.endswith("E-8-1")), ""),
                                    r"원문 범위로만[^|]*|범위 확인 필요[^|]*|미확정", rows=b.excel_rows_kesg61)),
            ("변형본 응답서 E-8-1을 확정 실적(verified·compared)으로 바꿈", "응답서 변형본 E-8-1",
             lambda b: _set_answer_in(b.sheet_kesg61, "E-8-1", {"status": "verified", "comparison": "compared"})),
            ("변형본 보고서 PDF 미확정 상태 제거", "보고서 PDF 변형본 E-8-1",
             lambda b: _pdf_sub(b, r"\[범위 미확정\][^.]*?(환경\s*법규\s*위반[^.]*?0(?:\.0)?\s*건)[^.]*\.", r"\1을 기록했습니다.")
             or _pdf_append(b, " FY2026 환경 법규 위반은 0건으로 확정되었습니다. ")),
        ]
    return cases


def _strip_marks(b: Bundle, pattern: str) -> int:
    b.body, n = re.subn(pattern, "", b.body)
    b.report_pdf_text, _ = re.subn(pattern.replace(r"\s*", r"\s*"), "", b.report_pdf_text)
    return n


def _pdf_append(b: Bundle, text: str) -> int:
    b.report_pdf_text += text
    return 1


def _pdf_sub(b: Bundle, pattern: str, repl: str, count: int = 1) -> int:
    at = b.report_pdf_text.find("지표 해설")
    head, tail = (b.report_pdf_text[:at], b.report_pdf_text[at:]) if at >= 0 else ("", b.report_pdf_text)
    tail, n = re.subn(pattern, repl, tail, count=count, flags=re.S)
    b.report_pdf_text = head + tail
    return n


def _sheet_pdf_cut(b: Bundle, text: str) -> int:
    flat_pattern = r"\s*".join(map(re.escape, glyph_flat(text)))
    b.sheet_pdf_text, n = re.subn(flat_pattern, "", b.sheet_pdf_text)
    return n


def _set_answer(b: Bundle, qid: str, key: str, value) -> int:
    for a in b.sheet["answers"]:
        if a["qid"] == qid:
            a[key] = value
            return 1
    return 0


def _excel_strip(b: Bundle, qid: str, pattern: str, rows: dict | None = None) -> int:
    rows = b.excel_rows if rows is None else rows
    row = rows.get(qid)
    if not row:
        return 0
    joined = " | ".join(row)
    stripped, n = re.subn(pattern, "", joined)
    rows[qid] = stripped.split(" | ")
    return n


def _set_answer_in(sheet: dict | None, suffix: str, values: dict) -> int:
    for a in (sheet or {}).get("answers", []):
        if a["qid"].endswith(suffix):
            a.update(values)
            return 1
    return 0


def _datasheet_strip(b: Bundle, code: str, column: str) -> int:
    row = b.datasheet.get(code)
    if not row or not row.get(column):
        return 0
    row[column] = ""
    return 1


# ── PR71 재검토 §8.3: 관계 오답(값·집단·날짜·사업장·사유)을 Markdown·PDF 양쪽(서로 일치)·한쪽에 넣는다 ─────────

def _md_add(b: Bundle, text: str) -> int:
    anchor = "### 주요 활동\n"
    if anchor not in b.body:
        return 0
    b.body = b.body.replace(anchor, anchor + text + "\n", 1)
    return 1


def _sheet_add(b: Bundle, text: str, where: str) -> int:
    """응답서 E-6(교육) 근거·Excel 행·응답서 PDF 중 한 곳에 문장을 더한다."""
    if where == "answer":
        a = next((a for a in b.sheet["answers"] if a["qid"].endswith("E-6")), None)
        if a is None:
            return 0
        a["rationale"] = (a.get("rationale") or "") + " " + text
        return 1
    if where == "excel":
        key = next((k for k in b.excel_rows if k.endswith("E-6")), None)
        if key is None:
            return 0
        b.excel_rows[key] = b.excel_rows[key] + [text]
        return 1
    if b.sheet_pdf_pages == 0:
        return 0
    b.sheet_pdf_text += "\n" + text + "\n"
    return 1


RELATION_ERRORS = [
    # (이름, 넣을 문장, Markdown 실패 항목 접두, 보고서 PDF 실패 항목 접두)
    ("원문에 없는 수량(정규직 15명)", "2026년 4월 22일 교육에는 정규직 15명과 기간제 6명이 출석했다.",
     "본문 교육 인원", "보고서 PDF 교육 인원"),
    ("원문에 없는 수량 표 행(999명)", "| 교육 참석 인원 | 999명 | 2026-04-22 |", "본문 교육 인원", "보고서 PDF 교육 인원"),
    ("40명·6명 고용형태 연결 뒤바꿈", "2026년 4월 22일 교육에 정규직 6명과 기간제 40명이 출석했다.",
     "본문 교육 인원 고용형태별", "보고서 PDF 교육 인원 고용형태별"),
    ("4/22 참석 46명을 4/27 참석으로", "2026년 4월 27일 교육에는 46명이 참석하였다.", "본문 교육 인원", "보고서 PDF 교육 인원"),
    ("다른 사업장 값을 해당 사업장 실적으로", "2026년 4월 22일 양산 제2공장 교육에는 46명이 참석하였다.",
     "본문 정답 수치를 다른 사업장", "보고서 PDF 정답 수치를 다른 사업장"),
]
RELATION_CONTROLS = [
    ("정상 정규직 40명·기간제 6명", "2026년 4월 22일 교육에 정규직 40명과 기간제 6명이 출석했다."),
    ("정상 순서만 바꿈(기간제 6명·정규직 40명)", "2026년 4월 22일 교육에 기간제 6명과 정규직 40명이 출석했다."),
    ("정상 같은 수의 두 집단(미참석 기간제 2명·파견 2명)", "2026년 4월 22일 교육 미참석 4명은 기간제 2명과 파견 2명이다."),
    ("정상 4/27 추가 참석 내역", "2026년 4월 27일 추가 교육에는 기간제 2명과 파견 2명이 참석했다."),
    ("정상 올바른 날짜의 46명", "2026년 4월 22일 교육 참석 인원은 46명이다."),
]


def relation_injections(*, variant: bool) -> tuple[list, list]:
    """(오답 [(이름, 실패해야 하는 항목 접두들, 변경 함수)], 정상 대조 [(이름, 변경 함수)])."""
    cases, controls = [], []
    for name, text, md_target, pdf_target in RELATION_ERRORS:
        md_text = text if text.startswith("|") else text
        pdf_text = text.strip("| ").replace(" | ", " ")
        cases += [
            (f"{name} — Markdown·PDF 양쪽(서로 일치)", [md_target, pdf_target],
             lambda b, m=md_text, t=pdf_text: _md_add(b, m) + _pdf_append(b, " " + t + " ")),
            (f"{name} — Markdown만", [md_target], lambda b, m=md_text: _md_add(b, m)),
            (f"{name} — 보고서 PDF만", [pdf_target], lambda b, t=pdf_text: _pdf_append(b, " " + t + " ")),
        ]
    swap = "2026년 4월 22일 교육에 정규직 6명과 기간제 40명이 출석했다."
    for where, label in (("answer", "응답서"), ("excel", "Excel"), ("sheet_pdf", "응답서 PDF")):
        cases.append((f"40명·6명 고용형태 연결 뒤바꿈 — {label}", [f"{label}의 교육 인원"],
                      lambda b, w=where: _sheet_add(b, swap, w)))
    cases.append(("교육 인원 오답 — 데이터시트", ["데이터시트 교육 인원"], lambda b: 0))   # N/A: 교육 인원 행이 없다
    if variant:
        cases += [
            ("변형본 Excel E-8-1 행에서 사유만 제거(상태 낱말은 남김)", ["Excel 변형본 E-8-1 행: 같은 행에 미확정 사유"],
             lambda b: _excel_strip(b, next((k for k in b.excel_rows_kesg61 if k.endswith("E-8-1")), ""),
                                    r"원문(?:에|의)[^|]*?확인하지\s*못했습니다\.?", rows=b.excel_rows_kesg61)),
            ("변형본 응답서 E-8-1에서 사유만 제거(상태 낱말은 남김)", ["응답서 변형본 E-8-1: 미확정 사유"],
             lambda b: _strip_answer_reason(b.sheet_kesg61, "E-8-1")),
        ]
    for name, text in RELATION_CONTROLS:
        controls += [(f"{name} — Markdown·PDF 양쪽", lambda b, t=text: _md_add(b, t) + _pdf_append(b, " " + t + " "))]
    for where, label in (("answer", "응답서"), ("excel", "Excel"), ("sheet_pdf", "응답서 PDF")):
        controls.append((f"정상 정규직 40명·기간제 6명 — {label}",
                         lambda b, w=where: _sheet_add(b, RELATION_CONTROLS[0][1], w)))
    return cases, controls


NOT_APPLICABLE = {"교육 인원 오답 — 데이터시트": "데이터시트는 K-ESG 프로필 항목만 싣고 교육 인원 행이 없다(원장 밖 원문 집계)"}


def _strip_answer_reason(sheet: dict | None, suffix: str) -> int:
    for a in (sheet or {}).get("answers", []):
        if a["qid"].endswith(suffix):
            n = 0
            for key in ("comparison_reason", "review_note"):
                text, k = re.subn(r"원문(?:에|의)[^|]*?확인하지\s*못했습니다\.?", "", str(a.get(key) or ""))
                a[key], n = text, n + k
            return n
    return 0


def run_injections(stage: Path, *, variant: bool) -> dict:
    base = load(stage)
    normal = check(copy.deepcopy(base), variant=variant)
    results = []
    cases = [(name, [target], mutate) for name, target, mutate in injections(base, variant=variant)]
    relation_cases, controls = relation_injections(variant=variant)
    for name, targets, mutate in cases + relation_cases:
        b = copy.deepcopy(base)
        changed = mutate(b)
        if not changed:
            results.append({"injection": name, "applied": False, "detected": None, "failed_items": [],
                            "na_reason": NOT_APPLICABLE.get(name, "주입할 자리가 없다")})
            continue
        r = check(b, variant=variant)
        failed = [row["item"] for row in r["rows"] if not row["pass"]]
        results.append({"injection": name, "applied": True, "target": targets,
                        "detected": all(any(item.startswith(t) for item in failed) for t in targets),
                        "failed_items": failed, "exit_code": 0 if r["passed"] == r["total"] else 1})
    control_results = []
    for name, mutate in controls:
        b = copy.deepcopy(base)
        changed = mutate(b)
        r = check(b, variant=variant) if changed else None
        control_results.append({"control": name, "applied": bool(changed),
                                "passed": None if r is None else r["passed"] == r["total"],
                                "failed_items": [] if r is None else [row["item"] for row in r["rows"] if not row["pass"]]})
    return {"stage": str(stage), "normal": {"passed": normal["passed"], "total": normal["total"],
                                            "failed": [r["item"] for r in normal["rows"] if not r["pass"]]},
            "injections": results, "controls": control_results,
            "all_detected": all(r["detected"] for r in results if r["applied"]),
            "controls_pass": all(c["passed"] for c in control_results if c["applied"]),
            "applied": sum(r["applied"] for r in results),
            "not_applicable": [{"injection": r["injection"], "reason": r["na_reason"]} for r in results
                               if not r["applied"]]}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("stage", type=Path)
    ap.add_argument("--json", type=Path)
    ap.add_argument("--variant", action="store_true", help="검증용 가상 변형본(14_) 포함 실행 — E-8-1 SOURCE_ONLY 항목 추가")
    ap.add_argument("--inject", action="store_true", help="오답 주입 사본으로 검사기 자체의 검출을 확인")
    args = ap.parse_args()
    if args.inject:
        result = run_injections(args.stage, variant=args.variant)
        for r in result["injections"]:
            state = "N/A" if not r["applied"] else ("DETECTED" if r["detected"] else "MISSED")
            print(f"{state} {r['injection']} :: exit={r.get('exit_code')} {r.get('failed_items', [])[:3]}"
                  + (f" ({r['na_reason']})" if not r["applied"] else ""))
        for c in result["controls"]:
            print(f"{'CONTROL-PASS' if c['passed'] else 'CONTROL-FAIL'} {c['control']} :: {c['failed_items'][:3]}")
        print(f"normal {result['normal']['passed']}/{result['normal']['total']} · injections applied {result['applied']}"
              f" · all detected {result['all_detected']} · controls pass {result['controls_pass']}")
        ok = result["all_detected"] and result["controls_pass"] and result["normal"]["passed"] == result["normal"]["total"]
    else:
        result = check(args.stage, variant=args.variant)
        for r in result["rows"]:
            print(f"{'PASS' if r['pass'] else 'FAIL'} [{r['where']}] {r['item']} :: {str(r['observed'])[:180]}")
        print(f"passed {result['passed']}/{result['total']} · {result['by_artifact']}")
        ok = result["passed"] == result["total"]
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
