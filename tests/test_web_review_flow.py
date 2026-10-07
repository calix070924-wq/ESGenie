"""Independent workflow contracts: replacement history, human edits and export parity."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import json

import fitz
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from esgenie.web.app import create_app
from esgenie.web.presenter import timestamp

HEADERS = {"X-ESGenie-Client": "workspace"}


def pdf(value):
    with fitz.open() as doc:
        doc.new_page().insert_text((40, 60), f"2026 January Factory A Electricity {value} kWh")
        return doc.tobytes()


def runner(project, directory):
    docs = [d for d in project["documents"] if d.get("included", True) and not d.get("error") and d["role"] == "evidence"]
    answers = []
    for index, doc in enumerate(docs):
        path = directory / doc["path"]
        with fitz.open(path) as source:
            quote = source[0].get_text().strip()
        value = float(quote.split()[-2])
        answers.append({"qid": f"q-{doc['id']}", "section": "환경", "question_text": doc["name"], "value": value,
                        "unit": "kWh", "status": "flagged", "boundary_label": "2026년 1월 · A공장 · 사용 전력",
                        "scope_notes": ["월간값 · 연간 실적 미확인"], "confidence_flags": ["partial_value"],
                        "evidence_links": [{"file_name": doc["name"], "quote": quote, "page": 0, "relative_path": "",
                                            "origin": "ocr_structured", "independent": True, "node_id": str(index)}]})
    return {"sheet": {"framework_key": "rba42", "framework_label": "RBA 참고양식", "answers": answers},
            "answers": [], "generated_at": timestamp(), "limitations": ["월간 자료만 있음"], "mode": "live"}


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(data_root=tmp_path, runner=runner, available=lambda: True), headers=HEADERS) as client:
        yield client


def setup(client):
    p = client.post('/api/projects', json={"company_name": "흐름 검증 회사", "year": 2026}).json()
    pid = p['id']
    for name, value in [('power.pdf', 142560), ('other.pdf', 400)]:
        p = client.post(f'/api/projects/{pid}/documents', files={'file': (name, pdf(value))}).json()
    return pid, p


def analyze(client, pid):
    assert client.post(f'/api/projects/{pid}/analysis').status_code == 202
    client.app.state.store.worker.submit(lambda: None).result(timeout=10)
    p = client.get(f'/api/projects/{pid}').json()
    assert p['job']['status'] == 'complete'
    return p


def save(client, pid, row, values=None, complete=False):
    response = client.put(f'/api/projects/{pid}/reviews/{row["id"]}', json={'values': values or row['saved'], 'complete': complete})
    assert response.status_code == 200, response.text
    return response.json()


def test_save_does_not_complete_and_screen_excel_pdf_share_human_values(client):
    pid, _ = setup(client)
    p = analyze(client, pid)
    row = p['result']['answers'][0]
    values = {**row['saved'], 'answer': '142,561.25', 'scope': '2026년 1월 · A공장 · 생산동', 'memo': '연간 자료 부족', 'reason': '원문 반올림 확인'}
    p = save(client, pid, row, values)
    shown = p['result']['answers'][0]
    assert shown['review_status'] == 'pending'
    assert shown['value_text'] == '142,561.25 kWh'
    assert shown['original_status'] == 'flagged'
    assert shown['history'][0]['before']['answer'] == '142,560'
    response = client.get(f'/api/projects/{pid}/download/bundle')
    assert response.status_code == 200, response.text
    with ZipFile(BytesIO(response.content)) as bundle:
        wb = load_workbook(BytesIO(bundle.read(next(n for n in bundle.namelist() if n.endswith('.xlsx')))))
        record = next(row for row in wb['응답서'].iter_rows(values_only=True) if row[0] == shown['id'])
        assert record[3] == shown['value_text']
        assert record[4] == shown['scope_label']
        assert record[7] == shown['review_label']
        assert record[8] == shown['method']
        assert record[9] == values['memo']
        assert values['sources'][0]['name'] in record[6]
        with fitz.open(stream=bundle.read('응답서.pdf'), filetype='pdf') as pdf_doc:
            text = ''.join(page.get_text() for page in pdf_doc)
        normalized = ''.join(text.split())
        for value in (shown['value_text'], shown['scope_label'], shown['review_label'], values['memo'], '원문 반올림 확인', 'power.pdf', '확인 필요'):
            assert ''.join(value.split()) in normalized
        audit = json.loads(bundle.read('검증_근거.json'))
        assert audit['sheet']['answers'][0]['value'] == 142560
    assert client.get(f'/api/projects/{pid}/download/pdf').content.startswith(b'%PDF')
    p = save(client, pid, shown, complete=True)
    assert p['result']['answers'][0]['review_status'] == 'complete'
    assert p['result']['answers'][0]['status'] == 'review'


def test_replacement_reanalysis_preserves_human_answer_other_completion_and_old_original(client):
    pid, _ = setup(client)
    p = analyze(client, pid)
    first, second = p['result']['answers']
    values = {**first['saved'], 'answer': '140,000', 'reason': '담당자 집계'}
    p = save(client, pid, first, values, complete=True)
    p = save(client, pid, second, complete=True)
    document = p['documents'][0]
    previous_pdf = client.get(f'/api/projects/{pid}/documents/{document["id"]}/original?version=1').content
    p = client.post(f'/api/projects/{pid}/documents/{document["id"]}/replace', files={'file': ('power-revised.pdf', pdf(150000))}).json()
    assert p['stale']
    assert p['documents'][0]['versions'][0]['name'] == 'power.pdf'
    assert p['affected_questions'] == [first['id']]
    assert p['result']['answers'][1]['review_status'] == 'complete'
    assert p['result']['answers'][0]['sources'][0]['version'] == 1
    for kind in ('xlsx', 'pdf', 'bundle'):
        assert client.get(f'/api/projects/{pid}/download/{kind}').status_code == 409
    p = analyze(client, pid)
    assert p['result']['answers'][0]['value_text'] == '140,000 kWh'
    assert p['result']['answers'][0]['automatic']['answer'] == '150,000'
    assert p['result']['answers'][0]['review_status'] == 'again'
    assert p['result']['answers'][1]['review_status'] == 'complete'
    assert client.get(f'/api/projects/{pid}/documents/{document["id"]}/original?version=1').content == previous_pdf
    assert client.get(f'/api/projects/{pid}/documents/{document["id"]}/pages/0?version=1').status_code == 200
    invalid = client.put(f'/api/projects/{pid}/reviews/{first["id"]}', json={'values': p['result']['answers'][0]['saved'], 'complete': True})
    assert invalid.status_code == 400
    newer = p['result']['answers'][0]['automatic']
    p = save(client, pid, p['result']['answers'][0], newer, complete=True)
    assert p['result']['answers'][0]['value_text'] == '150,000 kWh'
    assert p['result']['answers'][0]['sources'][0]['version'] == 2
    assert p['result']['answers'][0]['review_status'] == 'complete'
    assert not p['affected_questions']


@pytest.mark.parametrize('answer, expected', [('', '답변 없음'), ('   ', '답변 없음'), ('0', '0 kWh')])
@pytest.mark.parametrize('complete', [False, True])
def test_empty_numeric_answer_does_not_display_unit_in_screen_or_exports(client, answer, expected, complete):
    pid, _ = setup(client)
    p = analyze(client, pid)
    row = p['result']['answers'][0]
    values = {**row['saved'], 'answer': answer, 'memo': '전력 집계 자료 부족'}
    p = save(client, pid, row, values, complete=complete)
    shown = p['result']['answers'][0]
    assert shown['value_text'] == expected
    assert shown['saved']['answer'] == answer
    assert shown['saved']['unit'] == 'kWh'
    assert shown['review_status'] == ('complete' if complete else 'pending' if answer.strip() else 'missing')
    xlsx = client.get(f'/api/projects/{pid}/download/xlsx')
    assert xlsx.status_code == 200
    wb = load_workbook(BytesIO(xlsx.content))
    record = next(r for r in wb['응답서'].iter_rows(values_only=True) if r[0] == row['id'])
    assert record[3] == expected
    assert record[7] == shown['review_label']
    assert record[9] == values['memo']
    response = client.get(f'/api/projects/{pid}/download/pdf')
    assert response.status_code == 200
    with fitz.open(stream=response.content, filetype='pdf') as doc:
        text = ''.join(page.get_text() for page in doc)
    normalized = ''.join(text.split())
    assert ''.join(expected.split()) in normalized
    assert ''.join(values['memo'].split()) in normalized
    assert '답변없음kWh' not in normalized


def test_broken_encrypted_documents_remain_actionable_and_excluded_inputs_are_not_analyzed(client):
    pid, p = setup(client)
    broken = client.post(f'/api/projects/{pid}/documents', files={'file': ('broken.pdf', b'broken')}).json()['documents'][-1]
    assert broken['status'] == '실패' and '교체' in broken['error']
    with fitz.open(stream=pdf(20), filetype='pdf') as doc:
        protected = doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='secret')
    encrypted = client.post(f'/api/projects/{pid}/documents', files={'file': ('locked.pdf', protected)}).json()['documents'][-1]
    assert '암호' in encrypted['error']
    did = p['documents'][0]['id']
    client.patch(f'/api/projects/{pid}/documents/{did}', json={'included': False})
    p = analyze(client, pid)
    assert len(p['result']['answers']) == 1
    assert p['result']['answers'][0]['question'] == 'other.pdf'


def test_forged_evidence_rejected_and_empty_completion_requires_shortage_reason(client):
    p = client.post('/api/examples').json()
    row = p['result']['answers'][-1]
    response = client.put(f'/api/projects/{p["id"]}/reviews/{row["id"]}', json={'values': row['saved'], 'complete': True})
    assert response.status_code == 400
    values = {**row['saved'], 'memo': '윤리 규정과 운영 기록 미제공'}
    p = save(client, p['id'], row, values, complete=True)
    assert p['result']['answers'][-1]['review_status'] == 'complete'
    assert p['result']['answers'][-1]['status_label'] == '답변 작성 필요'
    bad = {**values, 'sources': [{'name': '../other.pdf', 'quote': 'invented', 'page': 0}]}
    assert client.put(f'/api/projects/{p["id"]}/reviews/{row["id"]}', json={'values': bad}).status_code == 400


def test_graph_candidates_preserve_actual_quote_unit_and_boundary():
    from types import SimpleNamespace
    from esgenie.web.engine import build_candidates
    from esgenie.supplychain import get_framework
    from esgenie.ssot.evidence_graph import EvidenceNode
    from esgenie.ssot.boundary import Boundary
    framework = get_framework('hmc')
    question = next(q for q in framework.questions if q.qtype == 'numeric' and q.kesg_codes)
    node = EvidenceNode('source', question.kesg_codes[0], 12.5, 'kWh', 2026, 'ocr',
                        raw_text='summary must not replace quote', quote='2026년 1월 A공장 사용 전력 12.5 kWh',
                        source_file='source.pdf', page=0, boundary=Boundary(period_text='2026-01', site='A공장'))
    output = SimpleNamespace(evidence_graph=SimpleNamespace(nodes={'source': node}))
    sheet = SimpleNamespace(framework_key='hmc', answers=[SimpleNamespace(qid=question.qid)])
    result = build_candidates(output, sheet, [{'id':'doc','name':'source.pdf','pages':1,'version':2}])
    candidate = result[question.qid][0]
    assert candidate['answer'] == '12.5' and candidate['unit'] == 'kWh'
    assert '2026-01' in candidate['scope'] and 'A공장' in candidate['scope']
    assert candidate['sources'][0]['quote'] == node.quote
    assert candidate['sources'][0]['version'] == 2


def test_manual_evidence_page_validated_and_analysis_snapshots_are_exported(client):
    pid, p = setup(client)
    p = analyze(client, pid)
    row = p['result']['answers'][0]
    other = p['documents'][1]
    source = {'document_id':other['id'],'name':other['name'],'version':1,'page':0,'quote':'user-supplied quote', 'manual': True}
    values = {**row['saved'], 'sources':[source]}
    p = save(client, pid, row, values)
    assert p['result']['answers'][0]['sources'][0]['quote'] == '담당자 직접 연결 · 해당 원문 페이지 확인 필요'
    p = client.post(f'/api/projects/{pid}/documents/{p["documents"][0]["id"]}/replace', files={'file': ('power.pdf', pdf(150000))}).json()
    p = analyze(client, pid)
    assert any(h['action'].startswith('재분석') for h in p['result']['answers'][0]['history'])
    response = client.get(f'/api/projects/{pid}/download/bundle')
    assert response.status_code == 200
    with ZipFile(BytesIO(response.content)) as bundle:
        assert any(n.startswith('이전_분석_기록/') for n in bundle.namelist())
        assert any(n.startswith(f'원본_버전/{other["id"]}/1/') for n in bundle.namelist())
