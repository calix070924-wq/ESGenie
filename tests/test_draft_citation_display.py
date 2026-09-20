"""AI 초안 인용 표기 — 2026-09-20 §5-3 회귀.

근거 문서
---------
  docs/작업지시서_HMC_응답서_검증정합성_개선_2026-09-20.md §5-3
  docs/reviews/HMC_출력_지적사항_검토_2026-09-20.md

고친 결함(HMC 응답서 실측)
  · 사용자용 본문·출처 목록에 내부 노드 ID가 그대로 찍혔다([LOCAL_TXT_0043]).
  · 페이지를 모르는 인용에 p.1이 붙었다((page or 0) + 1).

고정하는 계약
  1. 사용자용 본문 인용은 [1] 형태이고, 같은 번호의 출처 목록이 문서명 + 실제
     페이지를 준다. 번호는 안정적이다(출처 목록 순서 = 본문 번호).
  2. 내부 draft_text와 감사 JSON의 draft_citations[].node_id는 그대로 남는다 —
     표시를 바꿨을 뿐 추적성을 버리지 않는다.
  3. 괄호를 전부 벗기지 않는다. 출처 목록에 없는 ID는 출처를 꾸며내지 않고
     '[출처 미확인]'으로 두고 검토 사유를 남긴다.
  4. 페이지를 모르면 p.1을 붙이지 않는다. 0-기준 page=0은 p.1로(2로 밀지 않는다).

입력 구분: 손으로 만든 통제 실험(Answer 직접 구성). OCR/LLM 호출 없음.
"""
from __future__ import annotations

import json

from openpyxl import load_workbook

from esgenie.supplychain.exporters import export_response_sheet, export_response_sheet_pdf
from esgenie.supplychain.render import UNRESOLVED_MARK, draft_body, draft_lines, source_lines
from esgenie.supplychain.schema import Answer, ResponseSheet


def _answer(text: str, citations: list[dict]) -> Answer:
    return Answer("HMC-E-7", "경영시스템", "[E-7] 의사소통", None, "draft_ready",
                  draft_text=text, draft_citations=citations)


_CITATIONS = [
    {"node_id": "LOCAL_TXT_0043", "source_file": "노사협의회_운영규정.pdf", "page": 2,
     "text_preview": "분기별 노사협의회를 개최한다", "retrieval": "code_match"},
    {"node_id": "LOCAL_TXT_0051", "source_file": "사내공지_2026.pdf", "page": 0,
     "text_preview": "게시판에 공지한다", "retrieval": "bm25_fallback"},
]


class TestBodyUsesNumbersNotNodeIds:
    def test_internal_ids_become_stable_numbers(self):
        ans = _answer("분기별로 협의회를 연다 [LOCAL_TXT_0043]. 공지는 게시판에 [LOCAL_TXT_0051].",
                      _CITATIONS)
        body, notes = draft_body(ans)
        assert "[1]" in body and "[2]" in body
        assert "LOCAL_TXT" not in body
        assert notes == []

    def test_original_draft_text_is_untouched(self):
        ans = _answer("본문 [LOCAL_TXT_0043]", _CITATIONS)
        draft_body(ans)
        assert ans.draft_text == "본문 [LOCAL_TXT_0043]"
        assert ans.draft_citations[0]["node_id"] == "LOCAL_TXT_0043"

    def test_audit_json_keeps_the_node_id_and_adds_the_display_form(self):
        ans = _answer("본문 [LOCAL_TXT_0043]", _CITATIONS)
        payload = json.loads(json.dumps(ans.to_dict(), ensure_ascii=False))
        assert payload["draft_citations"][0]["node_id"] == "LOCAL_TXT_0043"
        assert payload["draft_text"] == "본문 [LOCAL_TXT_0043]"
        assert payload["draft_display"] == "본문 [1]"
        assert payload["draft_sources"][0].startswith("[1] 노사협의회_운영규정.pdf")


class TestSourceList:
    def test_numbers_document_names_and_real_pages(self):
        lines = source_lines(_answer("x", _CITATIONS))
        assert lines == ["[1] 노사협의회_운영규정.pdf p.3",
                         "[2] 사내공지_2026.pdf p.1"]

    def test_unknown_page_is_omitted_not_guessed_as_page_one(self):
        lines = source_lines(_answer("x", [
            {"node_id": "N1", "source_file": "의사소통절차서.pdf", "page": None}]))
        assert lines == ["[1] 의사소통절차서.pdf"]

    def test_unknown_document_name_is_said_out_loud(self):
        lines = source_lines(_answer("x", [{"node_id": "N1", "source_file": "", "page": 4}]))
        assert lines == ["[1] 문서명 미확인 p.5"]

    def test_repeated_node_id_shares_one_number(self):
        cits = [*_CITATIONS, dict(_CITATIONS[0])]
        assert len(source_lines(_answer("x", cits))) == 2


class TestUnresolvedCitations:
    def test_unlisted_internal_id_is_flagged_not_invented(self):
        ans = _answer("본문 [LOCAL_TXT_0099] 이다.", _CITATIONS)
        body, notes = draft_body(ans)
        assert UNRESOLVED_MARK in body
        assert "LOCAL_TXT_0099" not in body
        assert notes and "출처 미해소" in notes[0]

    def test_brackets_are_not_all_stripped(self):
        """인용 표시를 전부 벗기면 어느 문장이 증빙에 걸렸는지 알 수 없다."""
        body, _ = draft_body(_answer("A [LOCAL_TXT_0043] B [LOCAL_TXT_0099]", _CITATIONS))
        assert body.count("[") == 2

    def test_non_id_brackets_pass_through(self):
        body, notes = draft_body(_answer("에너지(E-4-1) [E-4-1] 구간 [2026-05] 기준", _CITATIONS))
        assert "[E-4-1]" in body and "[2026-05]" in body
        assert notes == []


class TestOutputsShareTheDisplay:
    def _sheet(self) -> ResponseSheet:
        ans = _answer("분기별 노사협의회를 운영한다 [LOCAL_TXT_0043].", _CITATIONS)
        return ResponseSheet("hmc", "현대차 협력사 ESG 사전점검 (RBA v8.0 매핑 참고양식)",
                             "한울정밀공업", [ans])

    def test_excel_shows_numbered_citation_and_no_internal_id(self, tmp_path):
        ws = load_workbook(export_response_sheet(self._sheet(), tmp_path))["응답서"]
        cell = next(ws.cell(row=r, column=4).value for r in range(5, ws.max_row + 1)
                    if ws.cell(row=r, column=1).value == "HMC-E-7")
        assert "[AI 초안 — 승인 전]" in cell
        assert "[1]" in cell and "노사협의회_운영규정.pdf p.3" in cell
        assert "LOCAL_TXT" not in cell

    def test_pdf_shows_the_same_lines(self, tmp_path):
        import fitz
        path = export_response_sheet_pdf(self._sheet(), tmp_path, embed_evidence=False)
        with fitz.open(path) as doc:
            text = "\n".join(p.get_text() for p in doc)
        # 좁은 열에서 파일명이 줄바꿈으로 쪼개지므로 공백을 제거하고 대조한다.
        flat = "".join(text.split())
        assert "노사협의회_운영규정.pdfp.3" in flat
        assert "출처:[1]" in flat
        assert "LOCAL_TXT" not in text

    def test_excel_and_pdf_use_the_same_source_of_truth(self):
        lines = draft_lines(self._sheet().answers[0])
        assert lines[0] == "[AI 초안 — 승인 전]"
        assert lines[1].endswith("[1].")
        assert lines[2].startswith("출처: [1] ")

    def test_actual_streamlit_draft_panel(self):
        """실제 Streamlit 위젯 렌더 — 화면도 같은 표기를 쓴다."""
        import sys

        from streamlit.testing.v1 import AppTest

        sys.modules.pop("esgenie.ui.tabs", None)
        app = AppTest.from_string("""
import streamlit as st
from esgenie.ui.tabs import _render_supplychain_drafts
_render_supplychain_drafts(st.session_state.answers)
""")
        ans = _answer("분기별 노사협의회를 운영한다 [LOCAL_TXT_0043]. 미확인 [LOCAL_TXT_0099].",
                      _CITATIONS)
        app.session_state["answers"] = [ans]
        app.run(timeout=30)
        assert not app.exception
        text = "\n".join(x.value for x in app.markdown)
        assert "운영한다 [1]" in text
        assert "[1] 노사협의회_운영규정.pdf p.3" in text
        assert UNRESOLVED_MARK in text
        assert "LOCAL_TXT_0099" not in text
        # 감사 추적용 node_id는 '근거 발췌' 안에 남는다.
        assert "node_id: LOCAL_TXT_0043" in text
        assert any("출처 미해소" in w.value for w in app.warning)
