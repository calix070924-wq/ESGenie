"""출력 표기 일관성 회귀 — 2026-09-20.

근거 문서
---------
  docs/작업지시서_HMC_응답서_검증정합성_개선_2026-09-20.md §5-1, §7
  docs/reviews/HMC_출력_지적사항_검토_2026-09-20.md

고친 결함(HMC 응답서 실측)
  · 경영시스템 영역 문항이 E-1 → E-10 → E-11 → E-12 → E-2 순으로 찍혔다(문자열 정렬).
  · Excel A2·PDF 표지에 AI초안(승인 대기)이 빠져 4분할 합이 100%에 못 미쳤다.
  · 0.513216 TJ가 화면·제출본에서 0.5 / 0으로 뭉개졌고, 142560 kWh에 천 단위 구분이 없었다.
  · 양식 제목이 현대차가 실제 송부한 양식처럼 읽혔다(실제로는 RBA v8.0 매핑 참고양식).

입력 구분: 손으로 만든 통제 실험(Answer·ResponseSheet 직접 구성) + 선언적 양식 정의.
OCR/LLM 호출 없음.
"""
from __future__ import annotations

import pytest
from openpyxl import load_workbook

from esgenie.supplychain.exporters import export_response_sheet, export_response_sheet_pdf
from esgenie.supplychain.frameworks import get_framework
from esgenie.supplychain.frameworks.hmc import natural_code_key
from esgenie.supplychain.render import FOUR_WAY, coverage_parts, page_label, summary_line
from esgenie.supplychain.responder import build_response_sheet
from esgenie.supplychain.schema import Answer, ResponseSheet, format_amount


class TestNaturalCodeOrder:
    def test_double_digit_codes_sort_after_single_digit(self):
        codes = ["E-10", "E-2", "E-1", "E-12", "E-9"]
        assert sorted(codes, key=natural_code_key) == ["E-1", "E-2", "E-9", "E-10", "E-12"]

    def test_hmc_management_area_runs_e1_to_e12_in_order(self):
        qids = [q.qid for q in get_framework("hmc").questions if q.section == "경영시스템"]
        assert qids == [f"HMC-E-{i}" for i in range(1, 13)]

    def test_sub_questions_stay_next_to_their_parent(self):
        """하위 문항 묶음은 부모 코드 바로 뒤에 붙는다 — 정렬이 묶음을 흩지 않는다."""
        qids = [q.qid for q in get_framework("hmc").questions if q.section == "환경"]
        parent = qids.index("HMC-C-4")
        assert qids[parent + 1].startswith("HMC-C-4-")
        assert qids[parent + 2].startswith("HMC-C-4-")
        assert qids[parent + 3] == "HMC-C-5"

    def test_hmc_label_says_reference_form_and_keeps_its_key(self):
        fw = get_framework("hmc")
        assert fw.key == "hmc"                       # 저장된 응답·설정 호환
        assert fw.label == "현대차 협력사 ESG 사전점검 (RBA v8.0 매핑 참고양식)"
        assert "실사 응답서" not in fw.label


class TestNumberFormatting:
    @pytest.mark.parametrize("value,expected", [
        (142560.0, "142,560"),        # 불필요한 .0 제거 + 천 단위 구분
        (142560, "142,560"),
        (0.513216, "0.513216"),       # 원장값은 표시 목적으로 반올림하지 않는다
        (0.873988, "0.873988"),
        (360772.0, "360,772"),
        (88.397, "88.397"),
        (10.6, "10.6"),
        (0.0, "0"),
    ])
    def test_amounts_keep_every_digit(self, value, expected):
        assert format_amount(value) == expected

    def test_display_value_never_collapses_a_small_ledger_value(self):
        ans = Answer("HMC-C-8-E-4-1", "환경", "총 에너지 사용량", 0.513216, "verified",
                     unit="TJ", period=2026)
        assert ans.display_value == "0.513216 TJ (2026년)"

    def test_display_value_groups_thousands(self):
        ans = Answer("x", "환경", "사용전력량", 142560.0, "verified", unit="kWh", period=2026)
        assert ans.display_value == "142,560 kWh (2026년)"

    def test_non_numeric_values_pass_through(self):
        """수치로 읽히지 않는 값을 조용히 0으로 바꾸지 않는다."""
        assert format_amount("미입력") == "미입력"
        assert format_amount(None) == "None"


class TestPageLabel:
    def test_unknown_page_is_not_reported_as_page_one(self):
        assert page_label(_Link(page=None)) == ""

    def test_first_page_is_page_one_not_two(self):
        assert page_label(_Link(page=0)) == "p.1"
        assert page_label(_Link(page=2)) == "p.3"


class _Link:
    def __init__(self, page):
        self.page = page
        self.file_name = "01_전기요금청구서_2026-05.pdf"
        self.bbox = None


def _hmc_sheet() -> ResponseSheet:
    """입력 없는 HMC 응답서 — 커버리지 4분할 골격을 고정한다."""
    return build_response_sheet("hmc", corp_name="한울정밀공업")


class TestFourWayHeader:
    def test_hmc_question_count_matches_the_expanded_clauses(self):
        """RBA 42조항 + 지표 수치행 6개 = 48문항.

        수치행: C-4(재활용률·배출량 2) + C-5(대기 배출량 1) + C-8(GHG·에너지·재생 3).
        C-5 분리(2026-09-20, §5-2) 전에는 47문항이었다 — 대기배출 관리체계 존재 여부와
        배출량 수치를 한 칸에서 묻던 문항을 C-4/C-8 패턴으로 나눈 결과다.
        """
        assert len(get_framework("hmc").questions) == 48

    def test_all_four_shares_are_reported_and_sum_to_one_hundred(self):
        """4분할은 상호배타 집계다 — 하나가 빠지면 합이 100%에 못 미친다.

        값 자체는 게이팅 결과에 따라 달라진다. 옛 비율(68.1/10.6/0.0/21.3)에 맞추려
        출력을 조작하지 않는다 — 여기서 고정하는 것은 '넷 다 보고되고 합이 맞는가'다.
        """
        sheet = _hmc_sheet()
        parts = coverage_parts(sheet)
        assert [label for label, _ in parts] == [label for label, _ in FOUR_WAY]
        assert sum(pct for _, pct in parts) == pytest.approx(100.0, abs=0.2)

    def test_summary_line_names_all_four_and_marks_the_overlap(self):
        line = summary_line(_hmc_sheet())
        for label, _ in FOUR_WAY:
            assert label in line
        assert "검토필요" in line and "중복 집계" in line

    def test_excel_and_pdf_share_the_header_sentence(self, tmp_path):
        sheet = _hmc_sheet()
        wb = load_workbook(export_response_sheet(sheet, tmp_path))
        a2 = wb["응답서"]["A2"].value
        assert a2 == summary_line(sheet)
        import fitz
        with fitz.open(export_response_sheet_pdf(sheet, tmp_path, embed_evidence=False)) as doc:
            text = "\n".join(p.get_text() for p in doc)
        for label, _ in FOUR_WAY:
            assert label in text


class TestScopeColumnReachesTheOutputs:
    def _sheet(self) -> ResponseSheet:
        ans = Answer("HMC-C-8-E-4-1", "환경", "총 에너지 사용량", 0.873988, "self_reported",
                     unit="TJ", period=2026,
                     completeness="partial",
                     boundary_label="2026-05 · 월간 · 사용전력량 + 도시가스 · 부분",
                     comparison="scope_unconfirmed",
                     comparison_reason="원장값이 부분값",
                     scope_notes=["E-4-1: 부분값 — 전사·연간 총량임이 입증되지 않아 검증 보류"])
        return ResponseSheet("hmc", get_framework("hmc").label, "한울정밀공업", [ans])

    def test_excel_has_a_scope_column_with_the_boundary(self, tmp_path):
        wb = load_workbook(export_response_sheet(self._sheet(), tmp_path))
        ws = wb["응답서"]
        header = [c.value for c in ws[4]]
        assert "측정 범위 / 검토" in header
        col = header.index("측정 범위 / 검토") + 1
        cell = next(ws.cell(row=r, column=col).value for r in range(5, ws.max_row + 1)
                    if ws.cell(row=r, column=1).value == "HMC-C-8-E-4-1")
        assert "월간" in cell and "범위 확인 필요" in cell

    def test_pdf_shows_the_value_at_full_precision_with_its_scope(self, tmp_path):
        import fitz
        path = export_response_sheet_pdf(self._sheet(), tmp_path, embed_evidence=False)
        with fitz.open(path) as doc:
            text = "\n".join(p.get_text() for p in doc)
        assert "0.873988" in text
        assert "측정 범위 / 검토" in text
        assert "범위 확인 필요" in text
