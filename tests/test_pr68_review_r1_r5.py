"""PR #68 검토 보완 R1~R5 회귀 (2026-09-29).

재현 입력은 검토 파일(repro.py)의 변형 표를 그대로 옮긴 픽스처다
(tests/fixtures/ocr_numeric_review_r1/review_cases.json). 모두 가상 입력이며 외부 API를
호출하지 않는다. R2·R4는 표 추출 → 근거 그래프 → 원장 선택 → DataPoint → 응답서
(→ Excel·PDF)까지 실제 함수를 거친다.
"""
from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from esgenie.ssot import ocr_router as R
from esgenie.ssot.ocr_router import ExtractedMetric, ExtractedTable, TableCell

FIXTURES = Path(__file__).parent / "fixtures" / "ocr_numeric_review_r1"
CASES = json.loads((FIXTURES / "review_cases.json").read_text(encoding="utf-8"))["cases"]


@pytest.fixture(autouse=True)
def _offline(monkeypatch):
    monkeypatch.setattr(R, "_get_openai_key", lambda: None)
    monkeypatch.setattr(R, "_get_anthropic_key", lambda: None)
    monkeypatch.setattr(R, "_get_upstage_key", lambda: None)


# ---- 헬퍼 ----------------------------------------------------------------------

def _run(tables_rows, doc_type, *, extra=(), name="review-variant.pdf", pages=None):
    """repro.py와 같은 방식 — 칸마다 좌표가 있는 표 객체 + 같은 칸 텍스트 토큰."""
    tokens = [{"text": s, "bbox": None, "page": 0} for s in extra]
    tables = []
    for t, rows in enumerate(tables_rows):
        page = pages[t] if pages else 0
        cells = []
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                box = [0.05 + c * 0.16, 0.2 + t * 0.2 + r * 0.03, 0.19 + c * 0.16, 0.21 + t * 0.2 + r * 0.03]
                tokens.append({"text": text, "bbox": box, "page": page})
                cells.append(TableCell(row_index=r, column_index=c, content=text, bbox=box, page=page))
        tables.append(ExtractedTable(table_id=f"t{t}", row_count=len(rows), column_count=max(map(len, rows)),
                                     cells=cells, source="upstage_dp", page=page))
    return R._tokens_to_extraction(tokens, doc_type=doc_type, file_path=name, engine="upstage_dp", tables=tables)


def _case(name, **kw):
    c = CASES[name]
    return _run(c["tables"], c["doc_type"], extra=c["extra"], **kw)


def _elec(rows, **kw):
    return _run([rows], "kepco_bill", **kw)


def _usage(ext, unit=None):
    return [m for m in ext.metrics if m.source_detail.get("role") == "usage" and (unit is None or m.unit == unit)]


def _review(ext):
    return [r["reason"] for r in (ext.router_meta.get("table_metrics") or {}).get("review", [])]


def _checks(ext, name):
    return [c for c in (ext.router_meta.get("table_metrics") or {}).get("checks", []) if c["name"] == name]


def _pipeline(exts, framework="rba42"):
    from esgenie.ssot.audit_trace import build_data_points
    from esgenie.ssot.detector_5axis import detect_d1_numeric
    from esgenie.ssot.evidence_graph import build_unified_graph
    from esgenie.ssot.ssot_pipeline import extract_with_ssot
    from esgenie.supplychain.responder import build_response_sheet
    g = build_unified_graph(None, exts, corp_code="TEST", corp_name="가상회사", report_year=2026)
    report = SimpleNamespace(source="ssot_local", corp_code="TEST", corp_name="가상회사", report_year=2026,
                             fiscal_year=2026, kesg_data={}, sections={}, raw_text="")
    result = extract_with_ssot(report, g, profile="sme")
    codes = [c for c in ("E-4-1", "E-3-1", "E-5-1", "E-6-1", "E-6-2") if g.resolved_facts.get(c)]
    scores, evaluations = {}, {}
    for code in codes:
        fact = g.resolved_facts[code]
        axis = detect_d1_numeric(f"{result.mapped[code]['name']} {fact.value}{fact.unit}", code, g)
        scores[code], evaluations[code] = axis.score, axis.evaluation
    points = build_data_points(g, scores, target_codes=codes, d1_evaluations=evaluations)
    sheet = build_response_sheet(framework, corp_name="가상회사", extraction=result, data_points=points)
    answers = {code: a for a in sheet.answers for code in codes if a.qid.endswith(f"-{code}")}
    return g, {p.kesg_code: p for p in points}, sheet, answers


MISMATCH_TEXTS = ("150,000 kWh", "142,560 kWh", "50,586", "48,210", "배율 60", "원측정값 상충")


# ---- R1. 지침 단위 환산 ----------------------------------------------------------

def test_r1_mwh_readings_computed_into_kwh_usage_column():
    ext = _case("computed_usage_unit_differs_from_index")
    [m] = _usage(ext)
    assert (m.value, m.unit) == (1000, "kWh")
    check = m.source_detail["index_check"]
    assert check["inputs"] == {"previous": 1, "current": 2, "multiplier": 1}
    assert check["input_units"]["previous"] == check["input_units"]["current"] == "MWh"
    assert check["computed"] == 1000 and check["computed_unit"] == "kWh"
    assert "배율" in check["formula"]


@pytest.mark.parametrize("explicit,status", [("1", "mismatch"), ("1,000", "match")])
def test_r1_explicit_kwh_checked_against_mwh_readings(explicit, status):
    ext = _elec([["전월지침(MWh)", "당월지침(MWh)", "배율", "사용량(kWh)"], ["1", "2", "1", explicit]])
    [m] = _usage(ext)
    assert m.source_detail["index_check"]["status"] == status
    assert m.source_detail["index_check"]["computed"] == 1000
    assert ("explicit_vs_index_mismatch" in _review(ext)) is (status == "mismatch")


def test_r1_mixed_kwh_and_mwh_readings_are_not_a_decrease():
    ext = _case("computed_mixed_energy_units")
    [m] = _usage(ext)
    assert "index_decreased" not in _review(ext)
    # 1,000 kWh → 2 MWh = 1 MWh(=1,000 kWh). 단위를 버리고 2 − 1,000으로 읽지 않는다.
    assert m.unit == "MWh" and m.value == pytest.approx(1.0)
    assert m.source_detail["index_check"]["input_units"] == {"previous": "kWh", "current": "MWh"}


@pytest.mark.parametrize("header,reason", [
    (["전월지침(kWh)", "당월지침", "배율"], "index_unit_missing"),
    (["전월지침(kWh)", "당월지침(m3)", "배율"], "index_unit_incompatible"),
])
def test_r1_missing_or_incompatible_reading_units_hold_the_calculation(header, reason):
    ext = _elec([header, ["1,000", "2,000", "1"]])
    assert _usage(ext) == []
    assert reason in _review(ext)


def test_r1_gas_volume_readings_never_become_mj():
    # 열량 칸이 비어 있으면 지침 차는 부피(m³)로만 남는다 — 발열량을 채워 MJ로 바꾸지 않는다.
    ext = _run([[["전월지침(m3)", "당월지침(m3)", "사용량(MJ)"], ["31,580", "40,000", "-"]]], "gas_bill")
    assert [m for m in ext.metrics if m.unit in ("MJ", "GJ", "TJ")] == []
    assert [(m.value, m.unit, m.kesg_code_guess) for m in ext.metrics] == [(8420, "m³", None)]


def test_r1_missing_multiplier_is_not_filled():
    ext = _elec([["전월지침(MWh)", "당월지침(MWh)", "사용량(kWh)"], ["1", "2", "-"]])
    assert _usage(ext) == []
    assert "multiplier_missing" in _review(ext)


def test_r1_zero_usage_across_units_is_zero():
    ext = _elec([["전월지침(kWh)", "당월지침(MWh)", "배율", "사용량(kWh)"], ["1,000", "1", "1", "-"]])
    [m] = _usage(ext)
    assert (m.value, m.unit) == (0, "kWh")


# ---- R2. 검산 불일치가 최종 답변·출력까지 ----------------------------------------------

def test_r2_meter_mismatch_reaches_answer_with_both_values_formula_and_reason():
    g, dps, sheet, answers = _pipeline([_case("explicit_meter_mismatch")])
    dp = dps["E-4-1"]
    assert dp.comparison == "mismatch" and dp.verification == "unverified"
    assert "source_conflict" in dp.confidence_flags
    ans = answers["E-4-1"]
    assert ans.status != "verified" and ans.comparison == "mismatch"
    for text in MISMATCH_TEXTS:
        assert text in ans.review_note, text
    # 범위 문구가 불일치 경고를 대신하지 않는다 — 불일치 사유가 맨 앞에 온다.
    assert ans.review_note.startswith("실제 불일치")


def test_r2_derived_emission_inherits_the_mismatch():
    _, dps, _, answers = _pipeline([_case("explicit_meter_mismatch")])
    dp = dps["E-3-1"]
    assert dp.comparison == "mismatch" and dp.verification == "unverified"
    assert all(t in answers["E-3-1"].review_note for t in MISMATCH_TEXTS)


def test_r2_matching_explicit_value_has_no_warning():
    ext = _elec([["이전 지침", "당월 지침", "계기 배율", "당월 전력 사용량"],
                 ["48,210", "50,586", "60", "142,560 kWh"]],
                extra=["사용 기간: 2026-04-01 ~ 2026-04-30", "사업장: 김해 제1공장"])
    _, dps, _, answers = _pipeline([ext])
    for code in ("E-4-1", "E-3-1"):
        assert dps[code].comparison != "mismatch"
        assert "source_conflict" not in dps[code].confidence_flags
        assert "원측정값 상충" not in answers[code].review_note


def test_r2_mismatch_does_not_flag_unrelated_documents():
    water = _run([[["전월지침(m3)", "당월지침(m3)", "사용량(m3)"], ["12,320", "13,000", "680"]]], "water_bill",
                 name="water.pdf", extra=["사용 기간: 2026-04-01 ~ 2026-04-30"])
    waste = _run([[["재활용", "합계"], ["5,400 kg", "18,400 kg"]]], "waste_ledger",
                 name="waste.pdf", extra=["위탁 기간: 2026-04-01 ~ 2026-04-30"])
    _, dps, _, answers = _pipeline([_case("explicit_meter_mismatch"), water, waste])
    assert dps["E-4-1"].comparison == "mismatch"
    for code in ("E-5-1", "E-6-1"):
        assert dps[code].comparison != "mismatch" and "source_conflict" not in dps[code].confidence_flags
        assert not any("원측정값 상충" in n for n in dps[code].scope_notes)
    assert "원측정값 상충" not in answers["E-6-1"].review_note


def test_r2_mismatch_is_visible_in_excel_and_pdf(tmp_path):
    from openpyxl import load_workbook
    from esgenie.supplychain.exporters import export_response_sheet, export_response_sheet_pdf
    _, _, sheet, _ = _pipeline([_case("explicit_meter_mismatch")])
    wb = load_workbook(export_response_sheet(sheet, tmp_path))
    cells = "\n".join(str(c.value) for row in wb["응답서"].iter_rows() for c in row if c.value is not None)
    for text in MISMATCH_TEXTS:
        assert text in cells, text
    fitz = pytest.importorskip("fitz")
    with fitz.open(export_response_sheet_pdf(sheet, tmp_path, embed_evidence=False)) as doc:
        text = "".join("".join(p.get_text().split()) for p in doc)
    for needle in ("150,000kWh", "142,560kWh", "원측정값상충"):
        assert needle in text, needle


# ---- R3. 가스 열량(MJ) 사용량 ------------------------------------------------------

def test_r3_gas_usage_header_in_mj_is_heat():
    [m] = _case("gas_MJ_usage_header").metrics
    assert (m.value, m.unit, m.kesg_code_guess) == (360772, "MJ", "E-4-1")
    assert m.source_detail["role"] == "heat" and m.source_detail["unit_source"] == "header"


def test_r3_gas_usage_with_mj_in_cell_is_heat():
    [m] = _run([[["가스사용량"], ["360,772 MJ"]]], "gas_bill").metrics
    assert (m.value, m.unit, m.kesg_code_guess) == (360772, "MJ", "E-4-1")


def test_r3_gas_volume_stays_volume():
    [m] = _run([[["사용량(m3)"], ["8,420"]]], "gas_bill").metrics
    assert (m.value, m.unit, m.kesg_code_guess) == (8420, "m³", None)


def test_r3_unit_less_gas_usage_gets_no_default_unit():
    ext = _run([[["가스사용량"], ["360,772"]]], "gas_bill")
    assert [m for m in ext.metrics if m.value == 360772] == []
    assert "unit_missing" in _review(ext)


def test_r3_empty_usage_column_does_not_claim_template_role():
    from esgenie.ssot.ocr_table_metrics import extract_table_metrics
    rows = [["가스사용량(MJ)", "비고"], ["-", "검침 예정"]]
    tables = _run([rows], "gas_bill")  # 경로 확인용 실행(예외 없음)
    assert tables.metrics == []
    res = extract_table_metrics([], [ExtractedTable(
        table_id="t0", row_count=2, column_count=2, source="upstage_dp", page=0,
        cells=[TableCell(row_index=r, column_index=c, content=t, bbox=[0.1, 0.1, 0.2, 0.2], page=0)
               for r, row in enumerate(rows) for c, t in enumerate(row)])], doc_type="gas_bill")
    assert res.claimed_roles == set()


def test_r3_volume_and_heat_of_same_gas_kept_without_double_count():
    ext = _run([[["사용량(m3)", "사용열량(MJ)"], ["8,420", "360,772"]]], "gas_bill",
               extra=["사용 기간: 2026-04-01 ~ 2026-04-30"])
    got = sorted((m.value, m.unit, m.kesg_code_guess) for m in ext.metrics)
    assert got == [(8420, "m³", None), (360772, "MJ", "E-4-1")]
    _, dps, _, _ = _pipeline([ext])
    assert dps["E-4-1"].value == pytest.approx(0.360772) and dps["E-4-1"].unit == "TJ"


# ---- R4. 행 라벨(사업장·기간) 보존과 같은 원문일 때만 중복 제거 --------------------------

def test_r4_single_row_tables_keep_site_labels():
    ms = _usage(_case("two_sites_equal_values"))
    assert sorted(m.source_detail["row_label"] for m in ms) == ["김해 제1공장", "양산 제2공장"]
    assert all(m.value == 1000 for m in ms)


def test_r4_single_row_tables_keep_month_labels_as_period():
    ms = _usage(_case("two_months_equal_values"))
    assert sorted(m.period for m in ms) == ["2026년 4월", "2026년 5월"]
    assert "conflicting_values" not in _review(_case("two_months_equal_values"))


def test_r4_multi_row_table_keeps_each_site():
    ext = _elec([["사업장", "사용량(kWh)"], ["김해 제1공장", "1,000"], ["양산 제2공장", "1,000"]])
    assert sorted(m.source_detail["row_label"] for m in _usage(ext)) == ["김해 제1공장", "양산 제2공장"]


def test_r4_same_value_on_different_pages_is_two_facts():
    ext = _run([[["사용량(kWh)"], ["1,000"]], [["사용량(kWh)"], ["1,000"]]], "kepco_bill", pages=[0, 1])
    assert sorted(m.page for m in _usage(ext)) == [0, 1]


def test_r4_same_source_cell_from_table_and_text_is_one_fact():
    ext = _elec([["사용량(kWh)"], ["1,000"]])
    assert len(_usage(ext)) == 1


def test_r4_same_scope_different_values_keeps_review():
    ext = _run([[["사업장", "사용량(kWh)"], ["김해 제1공장", "1,000"]],
                [["사업장", "사용량(kWh)"], ["김해 제1공장", "1,200"]]], "kepco_bill")
    assert "conflicting_values" in _review(ext)


@pytest.mark.parametrize("case,axis", [("two_sites_equal_values", "site"), ("two_months_equal_values", "period")])
def test_r4_final_answer_uses_one_scope_for_usage_and_emission(case, axis):
    g, dps, _, answers = _pipeline([_case(case)])
    e41, e31 = dps["E-4-1"], dps["E-3-1"]
    # 두 범위를 합산하지 않는다: 1,000 kWh = 0.0036 TJ 하나.
    assert e41.value == pytest.approx(0.0036)
    key = "site" if axis == "site" else "period_text"
    assert e41.boundary.get(key) and e41.boundary.get(key) == e31.boundary.get(key)
    assert " / " not in e41.boundary.get(key)


# ---- R5. 재활용률 표시 자릿수 검산 ----------------------------------------------------

def _rate_case(recycled, total, rate_text):
    rows = [["재활용", "합계"], [f"{recycled} kg", f"{total} kg"]]
    return _run([rows], "waste_ledger", extra=[f"폐기물 재활용률 {rate_text}"])


# 규칙: 재계산값을 원문 표시 소수 자릿수로 사사오입(ROUND_HALF_UP)한 값 == 원문 수치.
@pytest.mark.parametrize("recycled,total,text,status", [
    ("29.3", "100", "29%", "match"),
    ("29.3", "100", "29.0%", "mismatch"),
    ("29.3", "100", "29.3%", "match"),
    ("29.3", "100", "29.30%", "match"),
    ("5,400", "18,400", "29.3%", "match"),       # BM 29.3478%
    ("0", "100", "0%", "match"),
    ("100", "100", "100%", "match"),
    ("100", "100", "100.0%", "match"),
    ("0.04", "100", "0.0%", "match"),            # 0.04% → 0.0
    ("0.04", "100", "0%", "match"),
    ("0.05", "100", "0.1%", "match"),            # 경계: 0.05 → 0.1(사사오입)
    ("0.05", "100", "0.0%", "mismatch"),
    ("29.25", "100", "29.3%", "match"),          # 경계: 29.25 → 29.3
    ("29.25", "100", "29.2%", "mismatch"),
    ("29.5", "100", "30%", "match"),
])
def test_r5_rate_rounding_uses_source_display_decimals(recycled, total, text, status):
    ext = _rate_case(recycled, total, text)
    [check] = _checks(ext, "waste_recycling_rate")
    assert check["status"] == status
    assert check["reported_text"] == text
    assert "ROUND_HALF_UP" in check["rounding"]
    assert (ext.router_meta.get("hitl_required") is True) is (status == "mismatch")


def test_r5_source_rate_string_is_preserved():
    ext = _rate_case("29.3", "100", "29.30%")
    [rate] = [m for m in ext.metrics if m.kesg_code_guess == "E-6-2"]
    assert rate.source_detail["raw_text"] == "29.30%" and rate.source_detail["display_decimals"] == 2


def test_r5_unknown_precision_is_a_limitation_not_a_mismatch():
    from esgenie.ssot.ocr_table_metrics import TableMetricResult, check_recycling_rate
    res = TableMetricResult()
    for role, value in (("total", 100.0), ("recycled", 29.3)):
        m = ExtractedMetric(role, value, "kg", "", None)
        m.source_detail = {"role": role, "row_label": ""}
        res.metrics.append(m)
    rate = ExtractedMetric("재활용률", 29.0, "%", "", "E-6-2")   # LLM 등 원문 문자열 없음
    check_recycling_rate(res, [rate])
    [check] = res.checks
    assert check["status"] == "precision_unknown" and "자릿수" in check["note"]


def test_r5_rate_not_checked_across_different_row_scopes():
    ext = _run([[["사업장", "재활용", "합계"], ["김해 제1공장", "", "100 kg"], ["양산 제2공장", "29.3 kg", ""]]],
               "waste_ledger", extra=["폐기물 재활용률 29.0%"])
    assert _checks(ext, "waste_recycling_rate") == []
    assert "rate_check_scope_differs" in _review(ext)
