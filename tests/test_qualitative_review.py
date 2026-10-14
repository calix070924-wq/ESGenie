"""Whole-question conclusions must not be inferred from partial policy evidence."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO

import fitz
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from esgenie.ssot.evidence_graph import EvidenceGraph, merge_ocr_extraction
from esgenie.ssot.ocr_router import DocChannel, ExtractedClause, OcrExtraction
from esgenie.supplychain.frameworks import get_framework
from esgenie.supplychain.qualitative_review import REVIEW_REQUIRED_FLAG, require_qualitative_review
from esgenie.supplychain.responder import build_response_sheet
from esgenie.supplychain.schema import Answer
from esgenie.web.presenter import present_sheet
from esgenie.web.app import create_app
from esgenie.web.review import reconcile, review_rows


PARTIAL = [
    ('C-2', '스크랩 기록', '금속 스크랩을 공정 내부에 재투입한다.'),
    ('C-4', '폐기물 기록', '외부 위탁 폐기물의 중량을 월별로 기록한다.'),
    ('D-2', '거래 규정', '금품·부당한 편의 제공을 금지한다.'),
    ('D-6', '신고 절차', '근무자는 익명 신고할 수 있고 신고자에게 보복하지 않는다.'),
    ('D-8', '자료 보호', '근태 제공본에서 임직원 실명과 서명을 제외한다.'),
    ('E-6', '안전 교육 실시', '파견 근로자에게 안전 교육을 실시하였다.'),
]


def sheet_for(code, section, text):
    graph = EvidenceGraph('SYN', '검증기업')
    extraction = OcrExtraction(
        '운영자료.pdf', DocChannel.UNSTRUCTURED, 'policy_manual',
        clauses=[ExtractedClause(section=section, text=text, page=0, rba_code_guess=code)],
        raw_text='2024년 8월 / 대전 제3공장\n' + section + '\n' + text,
    )
    merge_ocr_extraction(graph, extraction, report_year=2024)
    return build_response_sheet('rba42', evidence_graph=graph)


@pytest.mark.parametrize('code,section,text', PARTIAL)
def test_partial_evidence_is_visible_without_a_whole_question_yes(code, section, text):
    sheet = sheet_for(code, section, text)
    answer = next(a for a in sheet.answers if a.qid == 'RBA-' + code)
    assert answer.value is None and answer.status == 'flagged'
    assert any(text in link.quote for link in answer.evidence_links)
    assert not answer.evidence_needed  # A review limit is not missing evidence.
    assert '담당자 확인 필요' in answer.review_note
    assert '응답: 예' not in answer.rationale
    shown = next(row for row in present_sheet(sheet.to_dict(), []) if row['id'] == answer.qid)
    assert shown['status'] == 'review' and shown['needs_attention']
    assert shown['value_text'] != '예'
    assert shown['sources']


@pytest.mark.parametrize('code,section,text', PARTIAL)
def test_review_guard_keeps_source_facts_claims_scope_and_conflicts(code, section, text):
    q = next(q for q in get_framework('rba42').questions if q.qid == 'RBA-' + code)
    a = next(a for a in sheet_for(code, section, text).answers if a.qid == q.qid)
    a = replace(a, value=True, status='flagged',
                self_reports=[{'yn': '예', 'text': '회사 작성 답변'}],
                flags=['별도 조항 검사 미충족'], comparison='mismatch',
                comparison_reason='독립 기록과 회사 답변 충돌')
    before = deepcopy(a.to_dict())
    require_qualitative_review(a, q)
    for field in ('evidence_links', 'reference_links', 'self_reports', 'flags',
                  'boundary', 'boundary_label', 'comparison', 'comparison_reason'):
        assert a.to_dict()[field] == before[field]
    assert a.value is None and a.status == 'flagged'


@pytest.mark.parametrize('code,section,text', PARTIAL)
@pytest.mark.parametrize('value,status', [(False, 'verified'), (False, 'flagged'),
                                         (None, 'insufficient'), (None, 'not_applicable')])
def test_negative_unresolved_and_excluded_answers_are_not_rewritten(code, section, text, value, status):
    q = next(q for q in get_framework('rba42').questions if q.qid == 'RBA-' + code)
    a = Answer(q.qid, q.section, q.text, value, status, rationale='기존 판단')
    before = deepcopy(a.to_dict())
    require_qualitative_review(a, q)
    assert a.to_dict() == before


def test_even_explicit_training_program_remains_reviewable_and_human_completion_works():
    text = ('관리자와 근로자를 대상으로 방침·절차·개선목표 이행과 법규 충족을 위한 '
            '교육 프로그램을 실시하였다.')
    sheet = sheet_for('E-6', '교육 실시', text)
    raw = next(a for a in sheet.answers if a.qid == 'RBA-E-6')
    assert raw.value is None and raw.status == 'flagged'
    assert any(text in e.quote for e in raw.evidence_links)
    project = {'framework': 'rba42', 'documents': [], 'result': {'sheet': sheet.to_dict()}}
    shown = present_sheet(sheet.to_dict(), [])
    before = next(r for r in review_rows(project, shown) if r['id'] == raw.qid)
    assert before['review_status'] == 'pending' and before['review_label'] == '확인 필요'
    values = {**before['saved'], 'answer': '예', 'reason': '대상과 교육 내용을 원문에서 확인'}
    project['reviews'] = {raw.qid: {'values': values, 'status': 'complete',
                                  'method': '담당자 수정', 'history': []}}
    after = next(r for r in review_rows(project, shown) if r['id'] == raw.qid)
    assert after['value_text'] == '예' and after['review_status'] == 'complete'
    assert after['automatic']['answer'] == ''
    assert raw.value is None and raw.status == 'flagged'


def test_reanalysis_requires_review_again_without_erasing_the_human_answer():
    sheet = sheet_for(*PARTIAL[0])
    current = {'sheet': sheet.to_dict()}
    previous = deepcopy(current)
    old = next(a for a in previous['sheet']['answers'] if a['qid'] == 'RBA-C-2')
    old.update(value=True, status='verified', rationale='응답: 예', scope_notes=[])
    project = {'reviews': {'RBA-C-2': {'status': 'complete', 'values': {'answer': '예'}}}}
    reconcile(project, previous, current)
    assert project['affected_questions'] == ['RBA-C-2']
    assert project['reviews']['RBA-C-2']['status'] == 'again'
    assert project['reviews']['RBA-C-2']['values']['answer'] == '예'


@pytest.mark.parametrize('qid,value,status,comparison', [
    ('RBA-C-4-E-6-1', 18.4, 'self_reported', 'scope_unconfirmed'),
    ('RBA-C-4-E-6-2', 29.3, 'flagged', 'not_comparable'),
    ('RBA-E-7', True, 'verified', ''),
])
def test_numeric_subquestions_and_communication_are_outside_review_limit(qid, value, status, comparison):
    q = next(q for q in get_framework('rba42').questions if q.qid == qid)
    a = Answer(q.qid, q.section, q.text, value, status, comparison=comparison)
    before = deepcopy(a.to_dict())
    require_qualitative_review(a, q)
    assert a.to_dict() == before


@pytest.mark.parametrize('code,section,text', PARTIAL)
@pytest.mark.parametrize('legacy', [False, True])
def test_api_save_reload_and_exports_keep_guarded_answers_pending(tmp_path, code, section, text, legacy):
    sheet = sheet_for(code, section, text)
    sheet.answers = [a for a in sheet.answers if a.qid == 'RBA-' + code]
    payload = sheet.to_dict()
    if legacy:
        payload['answers'][0]['confidence_flags'].remove(REVIEW_REQUIRED_FLAG)
    app = create_app(data_root=tmp_path, available=lambda: False)
    with TestClient(app, headers={'X-ESGenie-Client': 'workspace'}) as client:
        p = client.post('/api/projects', json={'company_name': '상태 검증', 'year': 2024}).json()
        pid = p['id']
        stored = app.state.store.read(pid)
        stored['result'] = {'sheet': payload, 'answers': [], 'mode': 'live', 'limitations': []}
        stored['result_revision'] = stored['input_revision']
        app.state.store.save(stored)

        def current():
            row = client.get(f'/api/projects/{pid}').json()['result']['answers'][0]
            assert row['review_status'] == 'pending' and row['review_label'] == '확인 필요'
            assert row['value_text'] == '답변 없음' and row['sources']
            return row

        row = current()
        if legacy:
            stored['reviews'] = {row['id']: {'values': deepcopy(row['saved']), 'status': 'missing',
                                            'history': [], 'method': '자동 입력'}}
            app.state.store.save(stored)
            row = current()  # Previously misclassified saves need no new analysis.
        values = {**row['saved'], 'memo': '기존 근거의 적용 범위를 확인 중'}
        response = client.put(f'/api/projects/{pid}/reviews/{row["id"]}',
                              json={'values': values, 'complete': False})
        assert response.status_code == 200, response.text
        row = current()
        assert app.state.store.read(pid)['reviews'][row['id']]['status'] == 'pending'
        assert row['saved']['memo'] == values['memo']

        excel = client.get(f'/api/projects/{pid}/download/xlsx')
        assert excel.status_code == 200
        book = load_workbook(BytesIO(excel.content))
        exported = next(r for r in book['응답서'].iter_rows(values_only=True) if r[0] == row['id'])
        assert exported[3] == '답변 없음' and exported[7] == '확인 필요'
        pdf = client.get(f'/api/projects/{pid}/download/pdf')
        assert pdf.status_code == 200
        with fitz.open(stream=pdf.content, filetype='pdf') as document:
            exported_text = ''.join(''.join(page.get_text().split()) for page in document)
        assert '확인필요' in exported_text and '자료필요' not in exported_text

        values.update(answer='예', reason='담당자가 전체 요건과 원문을 확인')
        response = client.put(f'/api/projects/{pid}/reviews/{row["id"]}',
                              json={'values': values, 'complete': True})
        assert response.status_code == 200
        completed = client.get(f'/api/projects/{pid}').json()['result']['answers'][0]
        assert completed['review_status'] == 'complete' and completed['value_text'] == '예'
        assert completed['original_status'] == 'flagged'


@pytest.mark.parametrize('state', ['complete', 'again'])
@pytest.mark.parametrize('legacy', [False, True])
def test_saved_completion_and_recheck_take_priority(state, legacy):
    sheet = sheet_for(*PARTIAL[0]).to_dict()
    if legacy:
        next(a for a in sheet['answers'] if a['qid'] == 'RBA-C-2')['confidence_flags'] = []
    project = {'framework': 'rba42', 'documents': [], 'result': {'sheet': sheet},
               'reviews': {'RBA-C-2': {'status': state, 'values': {'answer': '', 'unit': '',
                            'scope': '', 'sources': [], 'memo': '담당자 기록', 'reason': ''}, 'history': []}}}
    before = deepcopy(project)
    row = next(r for r in review_rows(project, present_sheet(sheet, [])) if r['id'] == 'RBA-C-2')
    assert row['review_status'] == state
    assert project == before  # Display projection must not mutate saved decisions.


@pytest.mark.parametrize('case', ['no_evidence', 'unrelated_flag', 'negative', 'numeric'])
def test_real_missing_answers_do_not_become_review_pending(tmp_path, case):
    raw = sheet_for(*PARTIAL[0]).to_dict()
    raw['answers'] = [a for a in raw['answers'] if a['qid'] == 'RBA-C-2']
    item = raw['answers'][0]
    if case == 'no_evidence':
        item.update(status='insufficient', evidence_links=[], evidence_needed=['원문 자료'])
    elif case == 'unrelated_flag':
        item.update(confidence_flags=[], rationale='다른 사유로 확인이 필요한 빈 답변')
    elif case == 'negative':
        item.update(value=False)
    else:
        item.update(qid='RBA-C-4-E-6-2', value=29.3, unit='%', comparison='not_comparable')
    app = create_app(data_root=tmp_path, available=lambda: False)
    with TestClient(app, headers={'X-ESGenie-Client': 'workspace'}) as client:
        pid = client.post('/api/projects', json={'company_name': '자료 부족 검증', 'year': 2024}).json()['id']
        p = app.state.store.read(pid)
        p['result'] = {'sheet': raw, 'answers': [], 'mode': 'live'}
        p['result_revision'] = p['input_revision']
        app.state.store.save(p)
        row = client.get(f'/api/projects/{pid}').json()['result']['answers'][0]
        if case in ('no_evidence', 'unrelated_flag'):
            assert row['review_status'] == 'missing'
        values = {**row['saved'], 'answer': '', 'memo': '답변할 자료를 확인 중'}
        response = client.put(f'/api/projects/{pid}/reviews/{row["id"]}',
                              json={'values': values, 'complete': False})
        assert response.status_code == 200
        row = response.json()['result']['answers'][0]
        assert row['review_status'] == 'missing' and row['review_label'] == '자료 필요'
        assert app.state.store.read(pid)['reviews'][row['id']]['status'] == 'missing'
