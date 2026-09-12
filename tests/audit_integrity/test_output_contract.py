"""Exercise actual Streamlit widgets and application exporters, without network."""
import json
import sys
from types import SimpleNamespace

import fitz
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from esgenie.dart_client import CompanyReport
from esgenie.pipeline import _apply_survey_answers, _collect_ocr_extractions
from esgenie.ssot.evidence_graph import build_unified_graph, TextNode
from esgenie.ssot.ssot_pipeline import extract_with_ssot
from esgenie.supplychain.responder import build_response_sheet
from esgenie.supplychain.exporters import export_response_sheet, export_response_sheet_pdf


def result_and_sheet(yn, note='', document=False):
    survey = {'E-1-1': {'yn':yn, 'text':note}}
    graph = build_unified_graph(None, _collect_ocr_extractions(None, survey_answers=survey),
                                corp_code='T', corp_name='검증용', report_year=2025)
    if document:
        graph.add_text_node(TextNode('policy', '환경방침', '환경방침을 수립하고 점검한다.',
                                    'E-1-1', 'policy.pdf', 2))
    ledger = extract_with_ssot(CompanyReport('T','검증용','',2025,{}, {}, [],'synthetic'), graph, profile='sme')
    _apply_survey_answers(ledger, survey)
    result = SimpleNamespace(extraction=ledger, evidence_graph=graph, v15_trace=None,
                             export_paths={}, corp_name='검증용')
    sheet = build_response_sheet('saq5_env', extraction=ledger, evidence_graph=graph)
    return result, sheet


@pytest.mark.parametrize('yn', ['예','아니오'])
@pytest.mark.parametrize('note', ['', '설문 메모'])
def test_survey_json_excel_pdf_and_actual_ui(tmp_path, yn, note):
    result, sheet = result_and_sheet(yn, note)
    answer = sheet.answers[0]
    payload = json.loads(json.dumps(sheet.to_dict(), ensure_ascii=False))['answers'][0]
    assert payload['value'] is (yn == '예') and payload['status'] == 'self_reported'
    assert payload['display_value'] == yn
    wb = load_workbook(export_response_sheet(sheet, tmp_path))
    row = next(r for r in wb['응답서'].iter_rows(values_only=True) if r[0] == answer.qid)
    assert row[3] == yn and '자가신고' in row[4]
    with fitz.open(export_response_sheet_pdf(sheet,tmp_path,embed_evidence=False)) as doc:
        text = '\n'.join(page.get_text() for page in doc)
    assert yn in text and '자가신고' in text
    sys.modules.pop('esgenie.ui.tabs', None)
    app = AppTest.from_string('''
import streamlit as st
from esgenie.ui.tabs import _render_supplychain_answer_detail
_render_supplychain_answer_detail(st.session_state.result, st.session_state.answer, question_map={})
''')
    app.session_state['result'] = result
    app.session_state['answer'] = answer
    app.run(timeout=30)
    assert not app.exception
    text = '\n'.join(x.value for x in app.markdown)
    assert f'**답변**: {yn}' in text and '**신뢰**: ⚠️ 자가신고' in text


def test_conflict_and_numeric_metadata_are_preserved_by_exporters(tmp_path):
    from esgenie.supplychain.schema import Answer
    _, sheet = result_and_sheet('아니오', '방침 없음', document=True)
    assert sheet.answers[0].value is False and sheet.answers[0].status == 'flagged'
    sheet.answers.append(Answer('energy','환경','에너지 사용량',248.5,'flagged',
                                unit='TJ',period=2025,flags=['부분값 검토']))
    wb = load_workbook(export_response_sheet(sheet,tmp_path))
    rows = {r[0]:r for r in wb['응답서'].iter_rows(values_only=True)}
    assert rows[sheet.answers[0].qid][3] == '아니오'
    assert '검토필요' in rows[sheet.answers[0].qid][4]
    assert rows['energy'][3] == '248.5 TJ (2025년)' and '부분값 검토' in rows['energy'][5]
    with fitz.open(export_response_sheet_pdf(sheet,tmp_path,embed_evidence=False)) as doc:
        text = '\n'.join(page.get_text() for page in doc)
    assert '아니오' in text and '검토필요' in text and '248.5 TJ (2025년)' in text


@pytest.mark.parametrize('all_axes', [True,False])
def test_actual_verify_ui_never_renders_abstention_as_pass(all_axes):
    from esgenie.schemas import AxisScore
    from esgenie.layer3_detect import _build_risk_vector
    ab = lambda: AxisScore(0,abstain=True,abstain_reason='no_evidence')
    rv = _build_risk_vector(ab() if all_axes else AxisScore(0),
        ab() if all_axes else AxisScore(0),ab(),ab() if all_axes else AxisScore(0))
    det = SimpleNamespace(risk_score=None if all_axes else 0, risk_vector=rv)
    step = SimpleNamespace(detection=det)
    verify = SimpleNamespace(final=step,steps=[step],final_score=det.risk_score,
        final_band=rv.evaluation_label,hitl_required=True,converged=False,iterations_used=0,
        final_text='검증용',metadata={'threshold':30})
    sys.modules.pop('esgenie.ui.tabs',None)
    app = AppTest.from_string('''
import streamlit as st
from esgenie.ui.tabs import render_verify_tab
render_verify_tab(st.session_state.result,'E','',show_header=False,show_final_text=False)
''')
    app.session_state['result'] = SimpleNamespace(sections={'E':verify},risk_rows=[])
    app.run(timeout=30)
    assert not app.exception and not app.success
    warnings = '\n'.join(x.value for x in app.warning)
    assert rv.evaluation_label in warnings and 'D3_semantic' in warnings
    assert '수렴 완료' not in '\n'.join(x.value for x in app.caption)


def test_content_change_invalidates_streamlit_response_cache(monkeypatch):
    import importlib
    sys.modules.pop('esgenie.ui.tabs',None)
    tabs = importlib.import_module('esgenie.ui.tabs')
    monkeypatch.setattr(tabs,'st',SimpleNamespace(session_state={}))
    calls=[]
    monkeypatch.setattr(tabs,'respond_from_pipeline',lambda *a,**kw: calls.append(kw) or len(calls))
    result,_ = result_and_sheet('예')
    assert tabs._get_cached_response_sheet(result,'saq5_env') == 1
    assert tabs._get_cached_response_sheet(result,'saq5_env') == 1
    result.extraction.mapped['E-1-1']['value'] = '[설문] 아니오'
    assert tabs._get_cached_response_sheet(result,'saq5_env') == 2
    assert tabs._get_cached_response_sheet(result,'saq5_env',supplier_claims={'E-6-2':42}) == 3
    assert tabs._get_cached_response_sheet(result,'saq5_env',supplier_claims={'E-6-2':92}) == 4


def test_existing_failed_clause_audit_blocks_verified_badge():
    result,sheet = result_and_sheet('예',document=True)
    checked = build_response_sheet('saq5_env',extraction=result.extraction,
        evidence_graph=result.evidence_graph,
        policy_audit=[{'kesg_code':'E-1-1','passed':False,
                       'findings':[{'status':'missing','description':'정기 검토 주기'}]}])
    assert sheet.answers[0].status == 'verified'
    assert checked.answers[0].status == 'flagged'
    assert any('정기 검토 주기' in flag for flag in checked.answers[0].flags)
