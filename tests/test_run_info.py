"""esgenie/run_info.py — 실행 출처 표시(B-2).

작업지시서 B §3 통과 조건 4: 세 판정 경계, dirty 감지, 키 값 미기록.
"""
from __future__ import annotations

import json
import subprocess

import pytest

from esgenie import run_info


# ── 판정 경계 3종 + 경계값 ───────────────────────────────────────────────────
def _llm(hits=0, misses=0, live_calls=0):
    return {"hits": hits, "misses": misses, "live_calls": live_calls}


def _ocr(hits=0, misses=0):
    return {"hits": hits, "misses": misses}


def test_classify_cache_replay():
    """LLM live_calls == 0 ∧ OCR misses == 0 ∧ Upstage 실요청 == 0 → 캐시 재생."""
    assert run_info.classify(_llm(hits=6, live_calls=0), _ocr(hits=10, misses=0),
                             upstage_live_requests=0, upstage_replay=True) == run_info.REPLAY


def test_classify_fresh():
    """LLM hits == 0 ∧ OCR hits == 0 ∧ Upstage 재생 아님 → 신규 처리."""
    assert run_info.classify(_llm(hits=0, misses=17, live_calls=17),
                             _ocr(hits=0, misses=10),
                             upstage_live_requests=4, upstage_replay=False) == run_info.NEW


def test_classify_mixed():
    """그 밖 → 혼합."""
    assert run_info.classify(_llm(hits=3, misses=5, live_calls=5), _ocr(hits=10, misses=1),
                             upstage_live_requests=4) == run_info.MIXED


def test_classify_all_zero_is_replay():
    """아무것도 부르지 않은 실행은 1번 규칙에 먼저 걸려 '캐시 재생'이다.

    2번도 만족하지만 작업지시서가 1번을 먼저 보라고 정했다 — 순서가 결과를 가른다.
    """
    assert run_info.classify(_llm(), _ocr(), upstage_live_requests=0) == run_info.REPLAY


def test_classify_llm_live_only():
    """LLM만 실호출하고 OCR은 전부 캐시 적중 → 혼합(1·2 모두 불만족)."""
    assert run_info.classify(_llm(hits=0, live_calls=3), _ocr(hits=5, misses=0),
                             upstage_live_requests=0) == run_info.MIXED


def test_classify_ocr_miss_only():
    """OCR만 미스이고 LLM은 캐시 적중만 → 혼합."""
    assert run_info.classify(_llm(hits=6, live_calls=0), _ocr(hits=2, misses=1),
                             upstage_live_requests=0) == run_info.MIXED


# ── Upstage 항을 더한 이유를 고정하는 경계 ───────────────────────────────────
def test_classify_upstage_live_with_all_caches_hit_is_mixed():
    """Upstage를 실제로 불렀는데 LLM·OCR이 전부 캐시 적중 → 혼합(캐시 재생이 아니다).

    `ocr_cache`는 Upstage 응답 캐시가 아니라 Upstage 결과를 입력으로 받는 VLM 보정
    LLM 응답 캐시다. 그래서 OCR 캐시가 전부 적중해도 Upstage는 매번 불린다
    (P1-0 followup 실측: OCR 적중 10·미스 1에 Upstage 요청 4건). 두 캐시만 보면
    이 실행이 '캐시 재생'으로 찍힌다 — 그걸 막는다.
    """
    assert run_info.classify(_llm(hits=6, live_calls=0), _ocr(hits=11, misses=0),
                             upstage_live_requests=4,
                             upstage_replay=False) == run_info.MIXED


def test_classify_upstage_replay_with_all_caches_miss_is_mixed():
    """Upstage 기록을 재생했는데 LLM·OCR이 전부 미스 → 혼합(신규 처리가 아니다).

    OCR 원본 추출을 새로 하지 않았으므로 '신규 AI 처리'라고 말할 수 없다.
    """
    assert run_info.classify(_llm(hits=0, misses=17, live_calls=17),
                             _ocr(hits=0, misses=10),
                             upstage_live_requests=0,
                             upstage_replay=True) == run_info.MIXED


def test_classify_unknown_upstage_never_replay():
    """Upstage 실요청 수를 모르면 '캐시 재생'으로 판정하지 않는다."""
    assert run_info.classify(_llm(hits=6, live_calls=0), _ocr(hits=10, misses=0),
                             upstage_live_requests=None) == run_info.MIXED


def test_classify_unknown_upstage_still_allows_fresh():
    """모르더라도 나머지 두 조건으로 '신규 처리' 판정은 허용한다."""
    assert run_info.classify(_llm(hits=0, misses=17, live_calls=17),
                             _ocr(hits=0, misses=10),
                             upstage_live_requests=None) == run_info.NEW


def test_upstage_recorded_in_processing():
    info = run_info.build_run_info(
        llm_stats_end={"mode": "on", "hits": 0, "misses": 3, "live_calls": 3},
        llm_stats_start={"mode": "on", "hits": 0, "misses": 0, "live_calls": 0},
        ocr_stats={"hits": 0, "misses": 2, "mode": "miss"},
        upstage_live_requests=4)
    assert info["processing"]["upstage"] == {"live_requests": 4, "replay": False,
                                            "counted": True}
    assert "Upstage 실요청 4" in run_info.summary_line(info)
    assert ("Upstage 실요청", "4") in run_info.rows(info)


def test_upstage_unknown_is_marked_everywhere():
    """요청 수를 모르면 요약 줄·Excel 행에도 '확인 못 함'으로 적는다."""
    info = run_info.build_run_info(
        llm_stats_end={"mode": "on", "hits": 0, "misses": 3, "live_calls": 3},
        llm_stats_start={"mode": "on", "hits": 0, "misses": 0, "live_calls": 0},
        ocr_stats={"hits": 0, "misses": 2, "mode": "miss"})
    assert info["processing"]["upstage"]["counted"] is False
    assert info["processing"]["upstage"]["live_requests"] is None
    assert "Upstage 실요청 확인 못 함" in run_info.summary_line(info)
    assert ("Upstage 실요청", "확인 못 함") in run_info.rows(info)


# ── 스냅샷 차이 ──────────────────────────────────────────────────────────────
def test_llm_numbers_use_snapshot_difference():
    """stats()는 프로세스 누적값이라 시작 스냅샷과의 차이를 쓴다."""
    info = run_info.build_run_info(
        llm_stats_end={"mode": "on", "hits": 10, "misses": 20, "live_calls": 20,
                       "successes": 20, "failures": 0},
        llm_stats_start={"mode": "on", "hits": 4, "misses": 3, "live_calls": 3,
                         "successes": 3, "failures": 0},
        ocr_stats={"hits": 0, "misses": 5, "mode": "miss"})
    llm = info["processing"]["llm"]
    assert (llm["hits"], llm["misses"], llm["live_calls"]) == (6, 17, 17)
    assert llm["from_snapshot"] is True


def test_llm_numbers_without_snapshot_are_marked():
    """스냅샷 없이 부르면 누적값을 쓰고 그 사실을 표시한다."""
    info = run_info.build_run_info(
        llm_stats_end={"mode": "on", "hits": 10, "misses": 20, "live_calls": 20},
        ocr_stats={"hits": 0, "misses": 5, "mode": "miss"})
    llm = info["processing"]["llm"]
    assert (llm["hits"], llm["live_calls"]) == (10, 20)
    assert llm["from_snapshot"] is False
    assert "누적" in run_info.summary_line(info)


def test_ocr_counted_from_extractions():
    """ocr_extractions를 주면 ocr_cache.summarize로 센다."""
    class _Ext:
        router_meta = {"ocr_cache_hits": 3, "ocr_cache_misses": 2, "ocr_cache": "hit,miss"}

    info = run_info.build_run_info(
        llm_stats_end={"mode": "on", "hits": 0, "misses": 1, "live_calls": 1},
        llm_stats_start={"mode": "on", "hits": 0, "misses": 0, "live_calls": 0},
        ocr_extractions=[_Ext(), _Ext()])
    assert info["processing"]["ocr"] == {"hits": 6, "misses": 4, "mode": "hit,miss"}


# ── 모델·경계 표기 ──────────────────────────────────────────────────────────
def test_ocr_model_is_upstage_not_llm():
    """models.ocr는 Upstage Document Parse다.

    ocr_cache.model_name()은 캐시 키용 LLM 모델명(= openai_model)이라
    그것을 'ocr'에 쓰면 'OCR 모델이 gpt-4.1-mini'라는 틀린 말이 된다.
    """
    from esgenie.config import SETTINGS
    from esgenie.ssot.ocr_router import UPSTAGE_DP_MODEL

    info = run_info.build_run_info(llm_stats_end={"mode": "on"}, ocr_stats={})
    assert info["models"]["ocr"] == UPSTAGE_DP_MODEL
    assert info["models"]["ocr_provider"] == "upstage_document_parse"
    assert info["models"]["llm"] == SETTINGS.openai_model
    assert info["models"]["ocr_vlm"] == SETTINGS.openai_model


def test_generated_at_is_kst():
    info = run_info.build_run_info(llm_stats_end={"mode": "on"}, ocr_stats={})
    assert info["generated_at"].endswith("+09:00")


def test_timings_optional():
    """PipelineOutput.timings가 없어도 동작한다(현재 코드에 그 필드가 없다)."""
    assert "timings" not in run_info.build_run_info(llm_stats_end={"mode": "on"}, ocr_stats={})
    info = run_info.build_run_info(llm_stats_end={"mode": "on"}, ocr_stats={},
                                   timings={"ocr": 1.5})
    assert info["timings"] == {"ocr": 1.5}


def test_upstage_replay_flag():
    info = run_info.build_run_info(llm_stats_end={"mode": "on"}, ocr_stats={},
                                   upstage_replay=True, upstage_live_requests=0)
    assert info["processing"]["upstage"]["replay"] is True
    assert info["processing"]["upstage_replay"] is True   # 이전 판 호환 자리
    assert "Upstage 기록 재생" in run_info.summary_line(info)
    assert ("Upstage 기록 재생", "예") in run_info.rows(info)


# ── dirty 감지 (임시 git 저장소) ─────────────────────────────────────────────
def _git(tmp_path, *args):
    return subprocess.run(["git", *args], cwd=tmp_path, capture_output=True,
                          text=True, encoding="utf-8", errors="replace")


@pytest.fixture
def git_repo(tmp_path):
    if _git(tmp_path, "init").returncode != 0:
        pytest.skip("git을 쓸 수 없는 환경")
    _git(tmp_path, "config", "user.email", "t@example.com")
    _git(tmp_path, "config", "user.name", "t")
    (tmp_path / "a.txt").write_text("1", encoding="utf-8")
    _git(tmp_path, "add", "a.txt")
    _git(tmp_path, "commit", "-m", "init")
    return tmp_path


def test_dirty_false_on_clean_repo(git_repo):
    state = run_info.code_state(git_repo)
    assert state["dirty"] is False
    assert len(state["commit_sha"]) == 40


def test_dirty_true_on_modified_repo(git_repo):
    (git_repo / "a.txt").write_text("2", encoding="utf-8")
    assert run_info.code_state(git_repo)["dirty"] is True


def test_dirty_true_on_untracked_file(git_repo):
    (git_repo / "new.txt").write_text("x", encoding="utf-8")
    assert run_info.code_state(git_repo)["dirty"] is True


def test_code_state_unknown_outside_repo(tmp_path):
    """git 저장소가 아니면 예외를 내지 않고 unknown/None으로 둔다."""
    outside = tmp_path / "not_a_repo"
    outside.mkdir()
    state = run_info.code_state(outside)
    assert state["commit_sha"] == "unknown"
    assert state["dirty"] is None
    assert run_info.short_sha(state) == "unknown"


def test_short_sha():
    assert run_info.short_sha({"commit_sha": "d8a0de8a434e4be9"}) == "d8a0de8"
    assert run_info.short_sha(None) == "unknown"


# ── 키 값 미기록 ─────────────────────────────────────────────────────────────
FAKE_OPENAI = "sk-FAKEKEY-must-not-appear-0123456789"
FAKE_ANTHROPIC = "sk-ant-FAKEKEY-must-not-appear-9876543210"
FAKE_UPSTAGE = "up-FAKEKEY-must-not-appear-abcdefghij"


@pytest.fixture
def fake_keys(monkeypatch):
    """가짜 키를 환경변수와 SETTINGS에 넣는다."""
    from esgenie.config import SETTINGS

    monkeypatch.setenv("OPENAI_API_KEY", FAKE_OPENAI)
    monkeypatch.setenv("ANTHROPIC_API_KEY", FAKE_ANTHROPIC)
    monkeypatch.setenv("UPSTAGE_API_KEY", FAKE_UPSTAGE)
    monkeypatch.setattr(SETTINGS, "openai_api_key", FAKE_OPENAI, raising=False)
    monkeypatch.setattr(SETTINGS, "anthropic_api_key", FAKE_ANTHROPIC, raising=False)
    return (FAKE_OPENAI, FAKE_ANTHROPIC, FAKE_UPSTAGE)


def test_keys_recorded_as_state_only(fake_keys):
    info = run_info.build_run_info(llm_stats_end={"mode": "on"}, ocr_stats={})
    assert info["keys"] == {"openai": "설정됨", "anthropic": "설정됨", "upstage": "설정됨"}
    blob = json.dumps(info, ensure_ascii=False)
    for key in fake_keys:
        assert key not in blob


def test_keys_absent_state(monkeypatch):
    from esgenie.config import SETTINGS

    monkeypatch.delenv("UPSTAGE_API_KEY", raising=False)
    monkeypatch.setattr(SETTINGS, "openai_api_key", None, raising=False)
    monkeypatch.setattr(SETTINGS, "anthropic_api_key", None, raising=False)
    info = run_info.build_run_info(llm_stats_end={"mode": "on"}, ocr_stats={})
    assert info["keys"] == {"openai": "없음", "anthropic": "없음", "upstage": "없음"}


def test_no_key_in_rendered_lines_and_rows(fake_keys):
    info = run_info.build_run_info(llm_stats_end={"mode": "on"}, ocr_stats={})
    blob = "\n".join([run_info.footer_line(info), run_info.summary_line(info),
                      *[f"{k}{v}" for k, v in run_info.rows(info)]])
    for key in fake_keys:
        assert key not in blob


def test_no_key_in_excel_and_pdf(fake_keys, tmp_path):
    """Excel 모든 셀과 PDF 추출 텍스트에 키 문자열이 없어야 한다."""
    import openpyxl
    import pymupdf

    from esgenie.supplychain.exporters.excel import export_response_sheet
    from esgenie.supplychain.exporters.pdf import export_response_sheet_pdf

    from .fixtures.make_run_info_baseline import sample_sheet

    info = run_info.build_run_info(
        llm_stats_end={"mode": "on", "hits": 0, "misses": 3, "live_calls": 3},
        llm_stats_start={"mode": "on", "hits": 0, "misses": 0, "live_calls": 0},
        ocr_stats={"hits": 0, "misses": 2, "mode": "miss"})
    out = tmp_path / "response_sheet"
    xlsx = export_response_sheet(sample_sheet(), out, run_info=info)
    pdf = export_response_sheet_pdf(sample_sheet(), out, embed_evidence=False, run_info=info)

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    cells = "\n".join(str(c.value) for name in wb.sheetnames
                      for row in wb[name].iter_rows() for c in row if c.value is not None)
    with pymupdf.open(pdf) as doc:
        text = "\n".join(page.get_text() for page in doc)

    assert "실행정보" in "".join(wb.sheetnames).replace(" ", "")
    for key in fake_keys:
        assert key not in cells, "Excel에 키 값이 실렸다"
        assert key not in text, "PDF에 키 값이 실렸다"
    # 설정 여부는 실려야 한다.
    assert "설정됨" in cells


# ── 출력물 반영 ─────────────────────────────────────────────────────────────
def test_excel_adds_only_new_sheet(tmp_path):
    """기존 시트는 그대로 두고 '실행정보' 시트만 더한다."""
    import openpyxl

    from esgenie.supplychain.exporters.excel import export_response_sheet

    from .fixtures.make_run_info_baseline import sample_sheet

    info = run_info.build_run_info(llm_stats_end={"mode": "on"}, ocr_stats={})
    plain = openpyxl.load_workbook(
        export_response_sheet(sample_sheet(), tmp_path / "plain"), data_only=True)
    stamped = openpyxl.load_workbook(
        export_response_sheet(sample_sheet(), tmp_path / "stamped", run_info=info),
        data_only=True)
    assert stamped.sheetnames == [*plain.sheetnames, "실행정보"]
    for name in plain.sheetnames:
        a = {c.coordinate: c.value for row in plain[name].iter_rows() for c in row}
        b = {c.coordinate: c.value for row in stamped[name].iter_rows() for c in row}
        assert a == b, f"기존 시트 '{name}'의 셀이 바뀌었다"


def test_pdf_footer_on_every_page_and_summary_on_first(tmp_path):
    import pymupdf

    from esgenie.supplychain.exporters.pdf import export_response_sheet_pdf

    from .fixtures.make_run_info_baseline import sample_sheet

    info = run_info.build_run_info(
        llm_stats_end={"mode": "on", "hits": 0, "misses": 3, "live_calls": 3},
        llm_stats_start={"mode": "on", "hits": 0, "misses": 0, "live_calls": 0},
        ocr_stats={"hits": 0, "misses": 2, "mode": "miss"})
    pdf = export_response_sheet_pdf(sample_sheet(), tmp_path, embed_evidence=False,
                                    run_info=info)
    short = run_info.short_sha(info)
    with pymupdf.open(pdf) as doc:
        pages = [page.get_text() for page in doc]
    for i, text in enumerate(pages):
        flat = text.replace(" ", "")
        assert short in flat, f"{i+1}쪽에 커밋 앞 7자리가 없다"
        assert "신규처리" in flat or "캐시재생" in flat or "혼합" in flat
    assert "LLM실호출" in pages[0].replace(" ", ""), "첫 페이지 요약 한 줄이 없다"


def test_stamp_lines_have_fixed_prefix():
    """바닥글·요약 줄은 고정 표식으로 시작한다.

    PDF 바닥글은 canvas에 직접 그려 추출 텍스트에 섞인다(응답표를 문항별로 읽는
    쪽이 표 내용과 가를 수 있어야 한다).
    """
    info = run_info.build_run_info(llm_stats_end={"mode": "on"}, ocr_stats={})
    assert run_info.footer_line(info).startswith(run_info.STAMP_PREFIX)
    assert run_info.summary_line(info).startswith(run_info.STAMP_PREFIX)


def test_none_returns_empty_lines():
    assert run_info.footer_line(None) == ""
    assert run_info.summary_line(None) == ""
    assert run_info.rows(None) == []
