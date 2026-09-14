"""Actual L4/L5, full app and application exporters consume D1 coverage."""
import copy
import json
import sys
from types import SimpleNamespace
import fitz
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from tests.d1_fixtures import CASES, verify_case, make_evidence
from esgenie.layer6_report import _block_esg, ReportDoc
from esgenie.exporters.report_pdf import export_report_pdf
from esgenie.layer3_judge import judge_risk_vector
from esgenie.schemas import AxisScore

@pytest.fixture(autouse=True)
def restore_app_logging():
    # app.py configures parent logger propagation; do not pollute later caplog tests.
    import logging
    logger = logging.getLogger('esgenie')
    level, propagate, handlers = logger.level, logger.propagate, list(logger.handlers)
    yield
    logger.setLevel(level)
    logger.propagate = propagate
    logger.handlers = handlers

@pytest.mark.parametrize('llm_judge', [False,True], ids=['rule','mock_judge'])
@pytest.mark.parametrize('name,sentence,score,compared,unverified', CASES, ids=[c[0] for c in CASES])
def test_final_verification_and_audit(name, sentence, score, compared, unverified, llm_judge):
    verify, trace, graph, _ = verify_case(sentence, llm_judge=llm_judge)
    axis = verify.final.detection.risk_vector.D1_numeric
    assert axis.score >= score
    assert axis.evaluation['compared_claims'] == compared
    assert axis.evaluation['unverified_claims'] == unverified
    assert verify.hitl_required == bool(score or unverified)
    assert verify.converged == (not score and not unverified)
    payload = json.loads(json.dumps(trace.to_dict(), ensure_ascii=False, allow_nan=False))
    assert payload['summary']['numeric_compared_claims'] == compared
    assert payload['summary']['numeric_unverified_claims'] == unverified
    for sentence_record in payload['sentences']:
        d1 = sentence_record['risk_vector']['D1_numeric']
        for claim in d1['evaluation']['claims']:
            assert sentence_record['sentence_text'][claim['start']:claim['end']] == claim['raw']
            if claim['status'] == 'compared':
                assert set(claim['evidence_ids']) <= set(graph.nodes)
                assert claim['evidence_value'] == graph.resolved_facts[claim['code']].value
        if score or unverified:
            assert sentence_record['hitl_status'] == 'HITL_REQUIRED'


def test_no_graph_cannot_use_legacy_zero_to_converge():
    verify, trace, *_ = verify_case('용수 사용량은 100톤이다.', missing_graph=True)
    assert not verify.converged and verify.hitl_required
    assert trace.sentences[0].risk_vector.D1_numeric.abstain_reason == 'no_evidence'


def test_text_coverage_includes_lower_risk_neighbor_sentence():
    sentence = '재생에너지 사용 비율은 92%였다. 용수 사용량은 100톤이다.'
    verify, trace, *_ = verify_case(sentence)
    rv = verify.final.detection.risk_vector
    assert rv.D1_numeric.score == 1
    assert rv.numeric_evaluation['status'] == 'partial'
    assert rv.numeric_evaluation['unverified_claims'] == 1
    for record in rv.numeric_evaluation['claims']:
        assert sentence[record['start']:record['end']] == record['raw']
    assert not rv.evaluation_complete and verify.hitl_required


def test_llm_opinion_cannot_clear_measured_mismatch_or_coverage():
    verify, *_ = verify_case(CASES[-1][1])
    rv = verify.final.detection.risk_vector
    fake = SimpleNamespace(complete=lambda **kw: SimpleNamespace(
        content=json.dumps({'D1_numeric':{'verdict':'false_positive','llm_score':0,'rationale':'수치가 맞다고 가정'}}),
        used_mock=True,meta={'model':'test-double'}))
    after = judge_risk_vector(CASES[-1][1], rv, fake)
    assert after.D1_numeric.score == 1
    assert after.D1_numeric.evaluation == rv.D1_numeric.evaluation
    assert not after.evaluation_complete

@pytest.mark.parametrize('case', [CASES[0],CASES[1],CASES[4],CASES[5],CASES[6],CASES[7]], ids=lambda c:c[0])
def test_report_pdf_keeps_numeric_coverage(tmp_path, case):
    name, sentence, score, compared, unverified = case
    verify, trace, *_ = verify_case(sentence)
    block = _block_esg(SimpleNamespace(sections={'E':verify}), 'E')
    doc = ReportDoc('D1검증','',2025,'2026-09-14',[block])
    with fitz.open(export_report_pdf(doc,tmp_path)) as pdf:
        text = ' '.join(' '.join(page.get_text().split()) for page in pdf)
    assert f'비교 {compared}건' in text and f'미검증 {unverified}건' in text
    assert sentence in text.replace('\n', ' ') or all(token in text for token in sentence.split())
    if unverified:
        assert '부분 평가' in text and '사람 검토 필요' in text
        assert ('D1 평가불가' if not compared else 'D1 부분 평가') in text


def test_ssot_json_excel_preserve_measured_value_and_unverified_state(tmp_path):
    from dataclasses import replace
    from esgenie.pipeline import _build_risk_rows
    from esgenie.ssot.audit_trace import build_data_points, build_audit_trace_v15
    from esgenie.ssot.excel_exporter import export_datasheet
    _, graph, _ = make_evidence()
    # Negative absolute performance is unsuitable even when the raw ledger contains it.
    graph.resolved_facts['E-3-1'] = replace(graph.resolved_facts['E-3-1'], value=-100)
    scores, rows = _build_risk_rows(graph,target_codes=['E-4-2','E-3-1'])
    points = build_data_points(graph,scores,target_codes=['E-4-2','E-3-1'],
        d1_evaluations={r['K-ESG 코드']:r['D1 평가'] for r in rows})
    outputs = export_datasheet(build_audit_trace_v15('D1','D1',points,[]),tmp_path)
    data = json.loads(open(outputs['audit_json']).read())['data_points']
    assert data[0]['value'] == 31 and data[0]['unit'] == '%'
    assert data[1]['d1_risk'] is None and data[1]['verification'] == 'unverified'
    assert data[1]['d1_evaluation']['reasons'] == {'invalid_number':1}
    wb = load_workbook(outputs['xlsx'])
    assert wb['DataSheet']['C2'].value == 31
    assert wb['DataSheet']['G3'].value is None
    assert wb['DataSheet']['I3'].value == '평가불가'
    assert wb['DataSheet']['K3'].value == 1

@pytest.fixture(scope='module')
def app_base():
    from esgenie.pipeline import run
    return run('005930',areas=['E'],export_outputs=False,save_traces=False)

@pytest.mark.parametrize('case', [CASES[0],CASES[1],CASES[4],CASES[5],CASES[6],CASES[7]], ids=lambda c:c[0])
def test_actual_app_d1_status(app_base, case, record_property):
    from pathlib import Path
    name, sentence, score, compared, unverified = case
    result = copy.deepcopy(app_base)
    verify, trace, *_ = verify_case(sentence)
    result.sections['E'] = verify
    result.audit_traces['E'] = trace
    sys.modules.pop('esgenie.ui.tabs',None)
    at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py'),default_timeout=60)
    at.session_state['result'] = result
    at.session_state['area_select'] = 'E'
    at.session_state['expert_mode'] = True
    at.run()
    assert not at.exception
    text = '\n'.join(b.value for b in list(at.markdown)+list(at.caption)+list(at.warning))
    record_property('ui_coverage', {'case':name, 'compared':compared, 'unverified':unverified,
        'review_required': '담당자 이관' in text, 'automatic_pass': '자동 통과' in text})
    assert f'비교 {compared}건' in text and f'미검증 {unverified}건' in text
    if score or unverified:
        assert '담당자 이관' in text
        assert '자동 통과' not in text and '재생성 없이 통과' not in text
    if unverified:
        assert ('D1 평가불가' if not compared else 'D1 부분 평가') in text


def test_saved_actual_document_ledger_ui_and_exports(tmp_path, app_base):
    from pathlib import Path
    from tests.d1_fixtures import real_excerpt_evidence
    from esgenie.pipeline import _build_risk_rows
    from esgenie.ssot.audit_trace import build_data_points, build_audit_trace_v15
    from esgenie.ssot.excel_exporter import export_datasheet
    _, graph, ledger, source = real_excerpt_evidence()
    codes = [m['code'] for m in source['metrics']]
    scores, rows = _build_risk_rows(graph,target_codes=codes)
    points = build_data_points(graph,scores,target_codes=codes,
        d1_evaluations={r['K-ESG 코드']:r['D1 평가'] for r in rows})
    trace = build_audit_trace_v15('MOBIS','현대모비스',points,[])
    outputs = export_datasheet(trace,tmp_path)
    wb = load_workbook(outputs['xlsx'])
    by_code = {m['code']:m for m in source['metrics']}
    # Existing selection labels the '매립 제로화' row as a component and does not
    # finalize it. D1 must report this as missing evidence; do not invent a total.
    from esgenie.layer3_detect import score_d1_numeric
    assert graph.resolved_facts['E-6-2'] is None
    assert score_d1_numeric('폐기물 재활용 비율은 92.9%였다.',graph).abstain_reason == 'no_evidence'
    assert len(points) == 1 and points[0].kesg_code == 'E-6-1'
    for index, point in enumerate(points,2):
        source_metric = by_code[point.kesg_code]
        assert point.value == source_metric['value'] == ledger.mapped[point.kesg_code]['value']
        assert point.period == int(source_metric['period']) == 2024
        assert point.evidence_files[0].page == source_metric['source_page_index'] == 69
        assert wb['DataSheet'].cell(index,3).value == point.value
        assert wb['DataSheet'].cell(index,4).value == point.unit
        assert wb['DataSheet'].cell(index,5).value == point.period
        assert point.d1_evaluation['compared_claims'] == 1
    result = copy.deepcopy(app_base)
    result.extraction, result.evidence_graph, result.v15_trace, result.risk_rows = ledger, graph, trace, rows
    sys.modules.pop('esgenie.ui.tabs',None)
    at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py'),default_timeout=60)
    at.session_state['result'] = result
    at.session_state['area_select'] = 'E'
    at.session_state['expert_mode'] = True
    at.run()
    assert not at.exception
    displayed = '\n'.join(str(frame.value) for frame in at.dataframe)
    assert '72463' in displayed or '72,463' in displayed
    assert '2024' in displayed

@pytest.mark.parametrize('case', [CASES[5],CASES[6],CASES[7]], ids=lambda c:c[0])
def test_calibration_and_evaluation_preserve_partial_claims(case):
    from esgenie.calibrate import _simulate_vector, BASELINE
    from esgenie.evaluate import _case_rows, abstain_coverage, with_abstain_ignored
    from esgenie.benchmark import _case_abstain_info
    verify,*_ = verify_case(case[1])
    rv = verify.final.detection.risk_vector
    record = {'id':case[0],'category':'D1','label':'greenwash' if case[2] else 'clean','axes':{}}
    for name in ['D1_numeric','D2_modifier','D3_semantic','D5_timeseries']:
        axis = getattr(rv,name)
        record['axes'][name] = dict(axis.to_dict(), rule_score=axis.score,
            judgeable=True,verdict='false_positive',llm_score=0)
    replay = _simulate_vector(record,trigger=.25,rule_weight=.4)
    assert replay.D1_numeric.evaluation == rv.D1_numeric.evaluation
    assert replay.D1_numeric.score == case[2] and not replay.evaluation_complete
    rows = _case_rows([record],BASELINE)
    assert rows[0]['evaluation_complete'] is False
    assert rows[0]['numeric_evaluation']['unverified_claims'] == 1
    assert _case_abstain_info(replay,bool(case[2]))[0] == (not case[2])
    assert abstain_coverage(rows)['abstains']['total'] == int(not case[2])
    assert not with_abstain_ignored(rows)[0]['evaluation_complete']
